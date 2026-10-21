# SPDX-License-Identifier: Apache-2.0
"""Independently verify two-dimensional semantic spans, atomic pagination and native DOCX vMerge.

Expected anchors, owner grids, row heights, page allocations and source UTF-16 are fixed here.
Java manifests only cross-check provenance against independently parsed final PDF geometry.
Every complete new PDF page is rendered by Poppler and pinned PDFium at 192 DPI. Existing
render thresholds are unchanged; no baseline is created or updated.
"""
from __future__ import annotations

import argparse
from importlib import metadata
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document as WordDocument
from pypdf import PdfReader
from verify_adr_form import render
from verify_advanced_pdf import difference, require
from verify_inline_images import image_boxes
from verify_multi_floats import pdf_geometry
from verify_production import actual_text, contact_sheet
from verify_semantic_table_flow import blue_ink
from verify_semantic_tables import annotations, check_characters, close, page_logical, table_borders

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "examples/build/row-spans"
OUTPUT = ROOT / "build/verification/row-spans"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
OD = "{urn:com.mirakyux.export4j:semantic:1}"
SVG = "{http://www.w3.org/2000/svg}"
TOLERANCE = .003


def expected_sources(name):
    """Tuples are physical row/column, columnSpan/rowSpan, original cell-local source."""
    if name == "groups":
        values = [(0, 0, 2, 2, "GROUP P=\ufffc N=\ufffc"), (0, 2, 2, 1, "HEADER-A"), (1, 2, 2, 1, "HEADER-B"),
                  (2, 0, 2, 2, "G01-A"), (2, 2, 1, 1, "G01-B"), (2, 3, 1, 2, "G01-C"), (3, 2, 1, 1, "G01-D"),
                  (4, 0, 4, 2, "G02-A1\nG02-A2")]
        for group in range(3, 7):
            row, prefix = 6 + (group - 3) * 3, f"G{group:02}-"
            values.extend([(row, 0, 1, 3, "\n".join(prefix + f"A{number}" for number in range(1, 4))),
                           (row, 1, 1, 2, prefix + "B1\n" + prefix + "B2"), (row, 2, 2, 1, prefix + "C"),
                           (row + 1, 2, 2, 1, prefix + "D"),
                           (row + 2, 1, 3, 1, prefix + "E" + (" P=\ufffc N=\ufffc L \ufffc" if group == 6 else ""))])
        return values
    if name == "columns":
        values = [(0, 0, 2, 2, "COL P=\ufffc N=\ufffc"), (0, 2, 1, 1, "HA"), (1, 2, 1, 1, "HB")]
        for group in range(1, 6):
            row = group * 2
            values.extend([(row, 0, 1, 2, f"K{group:02}-A"), (row, 1, 2, 1, f"K{group:02}-B"),
                           (row + 1, 1, 2, 1, f"K{group:02}-C" + (" P=\ufffc N=\ufffc L \ufffc" if group == 5 else ""))])
        return values
    return [(0, 0, 3, 2, "TPL P=\ufffc\nN=\ufffc"), (2, 0, 2, 2, "JAVA-TEMPLATE\nBOUND"),
            (2, 2, 1, 1, "XML-A"), (3, 2, 1, 1, "XML-B"), (4, 0, 3, 2, "ALL-TEMPLATE P=\ufffc\nN=\ufffc L \ufffc")]


def contract(name):
    if name == "groups":
        return {"pages": 3, "physicalRows": 18, "step": 18, "widths": [100, 140, 70, 70],
                "rowHeights": [24, 24, 24, 24, 21, 21, *([24] * 12)],
                "regions": [(0, 20, 20, list(range(9))), (1, 20, 20, [0, 1, *range(9, 15)]),
                            (2, 20, 20, [0, 1, 15, 16, 17])]}
    if name == "columns":
        return {"pages": 2, "physicalRows": 12, "step": 16, "widths": [58, 58, 58], "rowHeights": [22] * 12,
                "regions": [(0, 24, 124, [0, 1, 2, 3, 4, 5]), (0, 222, 124, [0, 1, 6, 7, 8, 9]),
                            (1, 24, 24, [0, 1, 10, 11])]}
    return {"pages": 1, "physicalRows": 6, "step": 18, "widths": [100, 140, 140],
            "rowHeights": [21, 21, 24, 24, 21, 21], "regions": [(0, 20, 20, list(range(6)))]}


def check_sources(name, expected):
    value = json.loads((FIXTURES / f"{name}-source.json").read_text(encoding="utf-8"))
    settings = contract(name)
    require((value["physicalRows"], value["physicalColumns"], value["headerRows"])
            == (settings["physicalRows"], len(settings["widths"]), 2), "physical grid/header counts differ")
    require(len(value["anchors"]) == len(expected), "covered slots/rows were turned into fake source anchors")
    owners = [[None] * len(settings["widths"]) for _ in range(settings["physicalRows"])]
    for actual, (row, column, width, height, source) in zip(value["anchors"], expected, strict=True):
        require((actual["row"], actual["column"], actual["columnSpan"], actual["rowSpan"], actual["source"])
                == (row, column, width, height, source), "physical origin, spans or bound source differs")
        require(actual["sourceUtf16Length"] == len(source.encode("utf-16-le")) // 2,
                "original cell-local UTF-16 length differs")
        objects = [index for index, character in enumerate(source) if character == "\ufffc"]
        image = objects[-1:] if row >= 2 and objects else []
        require(actual["imageIndices"] == image and actual["fieldIndices"] == [index for index in objects if index not in image],
                "bound field/image source offsets differ")
        for target_row in range(row, row + height):
            for target_column in range(column, column + width):
                require(owners[target_row][target_column] is None, "independent fixture grid contains an overlap")
                owners[target_row][target_column] = f"{row}:{column}"
    require(all(owner is not None for row in owners for owner in row) and value["owners"] == owners,
            "ownerAt lost incoming row-span coverage or filled the middle gap incorrectly")
    return value


def same_box(actual, expected):
    return actual["pageIndex"] == expected["pageIndex"] and all(
        abs(float(actual[key]) - float(expected[key])) < TOLERANCE for key in ("x", "y", "width", "height"))


def check_geometry(name, reader, expected, settings):
    manifest = json.loads((FIXTURES / f"{name}-results.json").read_text(encoding="utf-8"))
    require(manifest["pages"] == settings["pages"] and manifest["bodyPageOffset"] == 0,
            "final page count/contents offset differs")
    placements = manifest["placements"]
    require(len(placements) == len(settings["regions"]), "atomic table group region omitted/repeated")
    painted = table_borders(reader)
    _, marked, _ = pdf_geometry(reader, settings["step"])
    anchor_positions = {(value["page"], value["source"], round(value["x"], 3), round(value["y"], 3)) for value in marked}
    body_counts = {}
    last = None
    for index, (placement, (page, x, y, physical_rows)) in enumerate(zip(placements, settings["regions"], strict=True)):
        require(placement["pageIndex"] == page, "connected row-span group moved to an incorrect physical page")
        for key, value in (("x", x), ("y", y), ("width", sum(settings["widths"])),
                           ("height", sum(settings["rowHeights"][row] for row in physical_rows))):
            close(placement[key], value, "fixed row-span table geometry differs: " + key)
        if name == "columns":
            require((placement["fragmentIndex"], placement["fragmentCount"]) == (index, 3),
                    "nested semantic stream lost contiguous region numbering")
        origins, current = {}, y
        for row in physical_rows:
            origins[row] = current
            current += settings["rowHeights"][row]
        real_rows = [row for row in physical_rows if any(anchor[0] == row for anchor in expected)]
        require([row["row"] for row in placement["rows"]] == real_rows,
                "IR must retain original nonempty anchor rows and omit fully covered rows")
        for row in placement["rows"]:
            source_row = row["row"]
            require(row["header"] == (source_row < 2) and row["repeatedHeader"] == (source_row < 2 and index > 0)
                    and not row["continuedRow"], "merged header/group was mislabeled as a split row")
            source_cells = [anchor for anchor in expected if anchor[0] == source_row]
            require(len(row["cells"]) == len(source_cells), "IR contains fake incoming/covered anchors")
            for cell, (_, column, columns, rows, original) in zip(row["cells"], source_cells, strict=True):
                require((cell["column"], cell["columnSpan"], cell["rowSpan"], cell["firstLineIndex"])
                        == (column, columns, rows, 0), "original physical source origin/spans/firstLineIndex changed")
                require(all(covered in physical_rows for covered in range(source_row, source_row + rows)),
                        "a vertically connected group was split across regions")
                box = {"pageIndex": page, "x": x + sum(settings["widths"][:column]), "y": origins[source_row],
                       "width": sum(settings["widths"][column:column + columns]),
                       "height": sum(settings["rowHeights"][source_row:source_row + rows])}
                actual = {"pageIndex": page, **{key: cell[key] for key in ("x", "y", "width", "height")}}
                require(same_box(actual, box), "source anchor border does not cover its complete two-dimensional rectangle")
                require(sum(same_box(border, box) for border in painted) == 1,
                        "final PDF has no unique merged border at fixed expected coordinates")
                resolved = original.replace("\ufffc", str(page + 1), 1).replace("\ufffc", str(settings["pages"]), 1)
                originals, resolved_lines = original.split("\n"), resolved.split("\n")
                require(len(cell["lines"]) == len(originals), "merged cell was unexpectedly rewrapped or clipped")
                start = 0
                for line_index, (line, source, final) in enumerate(zip(cell["lines"], originals, resolved_lines, strict=True)):
                    end = start + len(source)
                    require(line["start"] == start and line["end"] in (end, end + int(line_index < len(originals) - 1))
                            and line["logical"] == final, "bound original UTF-16 or final field source differs")
                    close(line["x"], box["x"] + 3, "horizontal span repeated cell padding")
                    close(line["y"], box["y"] + 3 + line_index * settings["step"], "vertical span repeated padding or line step")
                    require(line["width"] <= box["width"] - 6 + TOLERANCE
                            and line["y"] + settings["step"] <= box["y"] + box["height"] - 3 + TOLERANCE,
                            "complete source line leaves its padded merged rectangle")
                    if "\ufffc" in source:
                        require((page, final, round(line["x"], 3), round(line["y"], 3)) in anchor_positions,
                                "manifest field/image source misses final independently painted logical anchor")
                    if " L " in source:
                        last = {"pageIndex": page, **line, "cell": box}
                    start = end + 1
                if source_row >= 2:
                    body_counts[source_row, column] = body_counts.get((source_row, column), 0) + 1
    require(body_counts == {(row, column): 1 for row, column, _, _, _ in expected if row >= 2},
            "real body anchors lost, repeated or moved into synthetic rows")
    require(last is not None, "final real source anchor lost linked media")
    expected_painted = sum(len([anchor for anchor in expected if anchor[0] in region[3]]) for region in settings["regions"])
    require(len(painted) == expected_painted + (3 if name == "columns" else 0),
            "covered rows/slots invented visible table borders")
    return placements, last, len(painted)


def check_links(name, reader, source):
    links, images = annotations(reader), image_boxes(reader)
    require(len(links) == 2 and len(images) == 1 and images[0]["pixels"] == [8, 8],
            "last real anchor image/links lost or repeated")
    require({value.get("uri") for value in links} == {f"https://example.com/{name}-text", f"https://example.com/{name}-image"},
            "last real anchor passive targets changed")
    image, page = images[0], source["pageIndex"]
    size = 10 if name == "columns" else 12
    require(image["page"] == page and image["width"] == image["height"] == size, "inline source image moved or changed size")
    picture = next(value for value in links if value["uri"].endswith("-image"))
    text = next(value for value in links if value["uri"].endswith("-text"))
    for key in ("x", "y", "width", "height"):
        close(picture[key], image[key], "image hyperlink differs from actual painted image: " + key)
    ink = blue_ink(reader)
    require(len(ink) == 1 and ink[0]["pageIndex"] == text["pageIndex"] == picture["pageIndex"] == page,
            "last anchor colored text/media moved to another page")
    close(text["x"], ink[0]["x"], "text hyperlink misses actual underline left")
    close(text["x"] + text["width"], ink[0]["right"], "text hyperlink misses actual underline right")
    close(text["y"], source["y"], "text hyperlink misses complete source line top")
    close(text["height"], source["height"], "text hyperlink misses fixed line height")
    box = source["cell"]
    require(box["x"] + 3 <= image["x"] < image["x"] + size <= box["x"] + box["width"] - 3 + TOLERANCE,
            "last-anchor image used a compact source index as physical column")
    return links, images


def check_fixture(name):
    path, settings, expected = FIXTURES / f"{name}.pdf", contract(name), expected_sources(name)
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == settings["pages"] and not reader.get_fields(), "fixed PDF page/static contract differs")
    for page in reader.pages:
        close(page.mediabox.width, 420, "page physical width differs")
        close(page.mediabox.height, 300, "page physical height differs")
    sources = check_sources(name, expected)
    placements, last, borders = check_geometry(name, reader, expected, settings)
    content = "\n".join((page.extract_text() or "") + "\n" + "\n".join(page_logical(page)) for page in reader.pages)
    expression = r"G\d{2}-[A-Z]\d?" if name == "groups" else r"K\d{2}-[A-Z]"
    if name in ("groups", "columns"):
        expected_labels = [value for row, _, _, _, text in expected if row >= 2 for value in re.findall(expression, text)]
        require(re.findall(expression, content) == expected_labels, "final PDF body source order/uniqueness differs")
    else:
        require(all(content.count(value) == 1 for value in ("JAVA-TEMPLATE", "BOUND", "XML-A", "XML-B", "ALL-TEMPLATE")),
                "static Java/XML bound source lost or repeated")
    if name == "columns":
        require(content.count("ROW-COLUMN-PREFIX") == 1, "short initial columns duplicated/lost prefix")
    links, images = check_links(name, reader, last)
    svg = ET.parse(FIXTURES / f"{name}-preview.svg").getroot()
    require(len(svg.findall(f"{SVG}svg")) == settings["pages"] and len(svg.findall(f".//{SVG}image")) == 1
            and not svg.findall(f".//{SVG}a"), "static SVG page/image contract differs")
    logical = [node.get("aria-label") for node in svg.iter(f"{SVG}g") if node.get("aria-label") is not None]
    marked = actual_text(reader)
    require([value for value in logical if value in set(marked)] == marked, "PDF/SVG final logical field order differs")
    return {"pages": settings["pages"], "sourceAnchors": len(sources["anchors"]), "physicalRows": settings["physicalRows"],
            "tablePresentations": len(placements), "paintedBorders": borders, "lastAnchor": last,
            "links": links, "images": images, "characters": check_characters(path, 20),
            "fixedGeometryOriginalUtf16OwnerGridAndAtomicGroupsChecked": True}


def check_package(path, private_metadata, consumer=False):
    with ZipFile(path) as package:
        main = ET.fromstring(package.read("word/document.xml"))
        require(not main.findall(f".//{W}altChunk"), "fixture contains unsupported nested content")
        if not private_metadata:
            require(not any(key.startswith(OD) for node in main.iter() for key in node.attrib),
                    "standard-only fixture retained private Export4j attributes")
        tables = main.findall(f".//{W}tbl")
        require(len(tables) == 1, "standard DOCX table missing/duplicated")
        table = tables[0]
        require([int(value.get(W + "w")) for value in table.findall(f"{W}tblGrid/{W}gridCol")]
                == ([1200, 2200, 2200] if consumer else [2000, 2800, 2800]),
                "native DOCX physical grid widths differ")
        physical = table.findall(f"{W}tr")
        require(len(physical) == 6 and len(table.findall(f"{W}tr/{W}trPr/{W}tblHeader")) == 2,
                "native DOCX lost covered physical rows or merged header group")
        recovered, active, continuation_count, rich_cells = [], {}, 0, 0
        for row_index, row in enumerate(physical):
            column = 0
            for cell in row.findall(f"{W}tc"):
                span = cell.find(f"{W}tcPr/{W}gridSpan")
                width = 1 if span is None else int(span.get(W + "val"))
                merge = cell.find(f"{W}tcPr/{W}vMerge")
                state = None if merge is None else merge.get(W + "val", "continue")
                if state == "continue":
                    require(column in active and active[column]["columnSpan"] == width,
                            "standard vMerge continuation has no matching restart rectangle")
                    owner = active[column]
                    require(owner["row"] + owner["rowSpan"] == row_index, "standard vMerge chain is not contiguous")
                    owner["rowSpan"] += 1
                    require(not cell.findall(f".//{W}t") and not cell.findall(f".//{W}fldSimple") and not cell.findall(f".//{W}drawing"),
                            "covered DOCX continuation duplicated source content")
                    continuation_count += 1
                else:
                    require(state in (None, "restart"), "nonstandard vMerge value")
                    if not consumer and (row_index, column) in ((2, 0), (2, 2), (3, 2)):
                        runs = [run for run in cell.findall(f".//{W}r") if run.find(f"{W}t") is not None]
                        require(runs and all(run.find(f"{W}rPr/{W}u") is not None
                                and run.find(f"{W}rPr/{W}u").get(W + "val") == "single" for run in runs),
                                "native DOCX bound rich-cell underline was lost")
                        rich_cells += 1
                    text = []
                    for node in cell.iter():
                        if node.tag == W + "t": text.append(node.text or "")
                        elif node.tag == W + "br": text.append("\n")
                        elif node.tag in (W + "fldSimple", W + "drawing"): text.append("\ufffc")
                    anchor = {"row": row_index, "column": column, "columnSpan": width, "rowSpan": 1, "source": "".join(text)}
                    recovered.append(anchor)
                    if state == "restart": active[column] = anchor
                    else: active.pop(column, None)
                column += width
            require(column == 3, "DOCX physical row does not cover all standard grid columns")
        expected = [(0, 0, 3, 2, "CON-ROWS\nCON-HEADER"), (2, 0, 1, 2, "CON-V"), (2, 1, 2, 1, "CON-UP"),
                    (3, 1, 2, 1, "P=\ufffc N=\ufffcLINK\ufffc"), (4, 0, 3, 2, "CON-COVERED\nCON-END")] if consumer else [
                    (row, column, columns, rows, "WORD HEADER\nCONTINUED" if row == 0 else source)
                    for row, column, columns, rows, source in expected_sources("templates")]
        require([(value["row"], value["column"], value["columnSpan"], value["rowSpan"], value["source"]) for value in recovered] == expected,
                "metadata-free standard vMerge parser recovered wrong original compact anchors")
        require(continuation_count == 3, "standard vMerge continuation count differs")
        fields = main.findall(f".//{W}fldSimple")
        require([field.get(W + "instr") for field in fields] == ["PAGE", "NUMPAGES"], "native DOCX final anchor fields differ")
        if private_metadata:
            require(all(float(field.get(OD + "width")) == (18 if consumer else 10) for field in fields),
                    "explicit native field reserved width differs")
        relationships = ET.fromstring(package.read("word/_rels/document.xml.rels"))
        targets = sorted(value.get("Target") for value in relationships if value.get("TargetMode") == "External")
        require(targets == (["https://example.com/rows", "https://example.com/rows-image"] if consumer else
                           ["https://example.com/templates-image", "https://example.com/templates-text"]),
                "standard DOCX final-anchor hyperlinks differ")
        require(sum(value.startswith("word/media/") for value in package.namelist()) == 1, "standard DOCX media lost/repeated")
    word = WordDocument(path)
    require(len(word.tables) == 1 and len(word.tables[0].rows) == 6, "independent python-docx physical table read failed")
    return {"physicalRows": 6, "sourceAnchors": len(recovered), "vMergeContinuations": continuation_count,
            "nativeRichCellDecorations": rich_cells,
            "privateMetadataPresent": private_metadata, "nativeFields": 2, "externalLinks": 2, "images": 1}


def check_consumer(directory):
    path = directory / "output/consumer-row-spans.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == 1 and not reader.get_fields(), "isolated caller row-span PDF contract differs")
    close(reader.pages[0].mediabox.width, 320, "caller physical page width differs")
    close(reader.pages[0].mediabox.height, 240, "caller physical page height differs")
    expected = [(20, 20, 280, 44), (20, 64, 60, 48), (80, 64, 220, 24), (80, 88, 220, 24), (20, 112, 280, 44)]
    borders = table_borders(reader)
    require(len(borders) == len(expected), "caller invented borders for covered rows/columns")
    for actual, (x, y, width, height) in zip(borders, expected, strict=True):
        for key, value in (("x", x), ("y", y), ("width", width), ("height", height)):
            close(actual[key], value, "caller fixed merged border geometry differs: " + key)
    content = (reader.pages[0].extract_text() or "") + "\n" + "\n".join(page_logical(reader.pages[0]))
    require(all(content.count(value) == 1 for value in ("CON-ROWS", "CON-HEADER", "CON-V", "CON-UP", "CON-COVERED", "CON-END"))
            and page_logical(reader.pages[0]) == ["P=1 N=1LINK\ufffc"], "caller source anchors/fields lost or repeated")
    _, marked, _ = pdf_geometry(reader, 20)
    require(len(marked) == 1, "caller marked final source count differs")
    close(marked[0]["x"], 82, "caller last real source physical column differs")
    close(marked[0]["y"], 90, "caller last real source line origin differs")
    links, images = annotations(reader), image_boxes(reader)
    require(len(links) == 2 and len(images) == 1 and images[0]["width"] == images[0]["height"] == 8,
            "caller last real source image/links differ")
    require({value.get("uri") for value in links} == {"https://example.com/rows", "https://example.com/rows-image"},
            "caller passive targets changed")
    image_link = next(value for value in links if value["uri"].endswith("-image"))
    for key in ("x", "y", "width", "height"):
        close(image_link[key], images[0][key], "caller image hyperlink differs from actual painted image")
    require(82 <= images[0]["x"] < images[0]["x"] + 8 <= 298, "caller media confused compact source and physical grid")
    package = check_package(directory / "output/consumer-row-spans.docx", True, consumer=True)
    return {"pages": 1, "paintedBorders": len(borders), "sourceAnchors": 5, "physicalRows": 6, "links": links,
            "images": images, "package": package, "characters": check_characters(path, 20)}


def render_checked(path, name, result, contacts):
    first, second = render(path, OUTPUT / name)
    require(len(first) == len(second) == result["pages"], "independent renderer omitted a complete page")
    metrics = [difference(a, b, 2) for a, b in zip(first, second, strict=True)]
    require(all(value["meanAbsoluteChannelError"] < 4 and value["changedPixelFraction"] < .06 for value in metrics),
            "independent renderers disagree")
    result["rendererDifferences"] = metrics
    contacts.extend((f"{name} page {index + 1}", page) for index, page in enumerate(first))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--consumer-build-directory", type=Path)
    args = parser.parse_args()
    versions = {name: metadata.version(name) for name in ("pypdf", "pypdfium2", "Pillow")}
    require(versions == {"pypdf": "6.10.0", "pypdfium2": "5.13.0", "Pillow": "12.3.0"}, "use the pinned project Python")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    report, contacts = {"tools": versions, "renderComparison": {"dpi": 192, "blurPixels": 2, "meanLimit": 4, "changedFractionLimit": .06}}, []
    for name in ("groups", "columns", "templates"):
        result = check_fixture(name)
        render_checked(FIXTURES / f"{name}.pdf", name, result, contacts)
        report[name] = result
    require((FIXTURES / "row-split-rejected.txt").read_text(encoding="utf-8") == "ROWSPAN_ROW_SPLIT_REJECTED\n",
            "explicit rowSpan/allowRowSplit rejection missing")
    require((FIXTURES / "docx-standard-read.txt").read_text(encoding="utf-8") == "STANDARD_VMERGE_STRICT_IMPORT_PASSED\n",
            "strict engine metadata-free native DOCX import missing")
    require(json.loads((FIXTURES / "bindings.json").read_text(encoding="utf-8"))
            == {"javaValueBlock": True, "xmlStaticTable": True, "coveredRows": True}, "static Java/XML two-dimensional bindings differ")
    report["package"] = check_package(FIXTURES / "spans.docx", True)
    report["standardPackage"] = check_package(FIXTURES / "spans-standard.docx", False)
    directory = args.consumer_build_directory
    if directory is None:
        pointer = ROOT / "build/verification/consumer-build-directory.txt"
        require(pointer.is_file(), "run isolated caller or pass --consumer-build-directory")
        directory = Path(pointer.read_text(encoding="utf-8").strip())
    if not directory.is_absolute(): directory = ROOT / directory
    result = check_consumer(directory)
    render_checked(directory / "output/consumer-row-spans.pdf", "consumer", result, contacts)
    report["consumerBuildDirectory"], report["consumer"] = str(directory), result
    require(len(contacts) == 7, "fixed sample complete-page budget differs")
    contact_sheet(contacts, OUTPUT)
    (OUTPUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"pages": len(contacts), "sourceGridAndFixedGeometry": "passed", "nativeVmergeWithoutMetadata": "passed"}, indent=2))


if __name__ == "__main__":
    main()
