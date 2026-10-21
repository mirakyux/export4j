# SPDX-License-Identifier: Apache-2.0
"""Independent new-capability fixtures. Run after :examples:capabilities; never updates visual baselines."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from pypdf import PdfReader
from docx import Document
import pypdfium2 as pdfium
from verify_advanced_pdf import require, render_poppler, render_pdfium, difference
from verify_production import actual_text, extracted, contact_sheet

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "examples/build/capabilities"


def render_encrypted(path, output):
    poppler_dir, pdfium_dir = output / "poppler", output / "pdfium"
    poppler_dir.mkdir(parents=True, exist_ok=True)
    pdfium_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(["pdftoppm", "-upw", "fixture-reader", "-r", "96", "-png", str(path), str(poppler_dir / "page")],
                   check=True, capture_output=True, timeout=60)
    with pdfium.PdfDocument(path, password="fixture-reader") as document:
        require(len(document) == 1, "encrypted page count differs")
        page = document[0]
        try:
            bitmap = page.render(scale=96 / 72, draw_annots=True)
            try:
                bitmap.to_pil().convert("RGB").save(pdfium_dir / "page-1.png")
            finally:
                bitmap.close()
        finally:
            page.close()
    return poppler_dir / "page-1.png", pdfium_dir / "page-1.png"


def verify_signature(path, openssl, output):
    reader = PdfReader(path)
    fields = reader.trailer["/Root"]["/AcroForm"]["/Fields"]
    signatures = [field.get_object()["/V"] for field in fields if field.get_object().get("/FT") == "/Sig"]
    require(len(signatures) == 1, "expected one signature")
    signature = signatures[0]
    ranges = list(signature["/ByteRange"])
    raw = path.read_bytes()
    require(ranges[0] == 0 and ranges[1] < ranges[2] and ranges[2] + ranges[3] == len(raw), "signature byte range is incomplete")
    require(raw.startswith((FIXTURES / "invoice.pdf").read_bytes()), "signature changed the original revision")
    payload = raw[:ranges[1]] + raw[ranges[2]:ranges[2] + ranges[3]]
    cms = bytes(signature["/Contents"])
    require(cms[0] == 0x30, "CMS is not DER sequence")
    if cms[1] < 128:
        end = 2 + cms[1]
    else:
        length_bytes = cms[1] & 127
        require(0 < length_bytes <= 4, "invalid DER length")
        end = 2 + length_bytes + int.from_bytes(cms[2:2 + length_bytes], "big")
    require(end <= len(cms) and not any(cms[end:]), "invalid CMS padding")
    (output / "signature.der").write_bytes(cms[:end])
    data = output / "signed-content.bin"
    data.write_bytes(payload)
    command = [str(openssl), "cms", "-verify", "-binary", "-inform", "DER", "-in", str(output / "signature.der"),
               "-content", str(data), "-noverify", "-out", str(output / "verified-content.bin")]
    result = subprocess.run(command, capture_output=True, timeout=30)
    require(result.returncode == 0, "detached CMS cryptographic verification failed")
    require((output / "verified-content.bin").read_bytes() == payload, "verified signature content differs")
    tampered = bytearray(payload)
    tampered[len(tampered) // 2] ^= 1
    data.write_bytes(tampered)
    rejected = subprocess.run(command, capture_output=True, timeout=30)
    require(rejected.returncode != 0, "tampered signed content was accepted")
    data.write_bytes(payload)
    return {"byte_range": ranges, "cms_bytes": end, "cryptographic_signature": "passed", "tampering": "rejected", "certificate_trust": "self-signed development certificate; not validated"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--openssl", type=Path)
    parser.add_argument("--verapdf-report", type=Path)
    args = parser.parse_args()
    output = ROOT / "build/verification/capabilities"
    output.mkdir(parents=True, exist_ok=True)
    report = {}
    contacts = []
    for name in ("shared", "balanced", "invoice", "cff", "archive"):
        path = FIXTURES / f"{name}.pdf"
        reader = PdfReader(path)
        text = extracted(path)
        logical = "\n".join(actual_text(reader))
        (output / f"{name}-pdfium.txt").write_text(text, encoding="utf-8")
        (output / f"{name}-actualtext.txt").write_text(logical, encoding="utf-8")
        if name == "shared":
            for marker in ("BODY-01", "CELL-02", "HEADER-01", "NOTE-01"):
                require(marker in logical, f"shared shaping missed {marker}")
            require("a\u0301" in logical and "العربية لا" in logical, "shared logical text changed")
            links = [a.get_object() for page in reader.pages for a in page.get("/Annots", [])]
            require(len(links) == 1 and links[0]["/A"]["/URI"] == "https://example.com/docs?a=1&b=2", "external link destination or box count differs")
        elif name == "balanced":
            require(re.findall(r"BAL-(\d{2})", logical) == [f"{i:02d}" for i in range(1, 12)], "balanced columns lost/duplicated/reordered source")
            require("AFTER-COLUMNS" in text and len(reader.pages) == 1, "balanced final page did not leave room for following text")
        elif name == "invoice":
            require("SKU-001" in text and "SKU-002" in text and "已确认" in text, "template content differs")
        elif name == "cff":
            require("CFF-01" in logical and "a\u0301" in logical, "CFF logical text differs")
        elif name == "archive":
            require(reader.trailer["/Root"].get("/OutputIntents") and reader.trailer["/Root"].get("/Metadata"), "archival resources missing")
        poppler = render_poppler(path, output / name / "poppler")
        pdfium = render_pdfium(path, output / name / "pdfium")
        require(len(poppler) == len(pdfium) == len(reader.pages), "renderer page counts differ")
        metrics = [difference(a, b, 1) for a, b in zip(poppler, pdfium)]
        for metric in metrics:
            require(metric["meanAbsoluteChannelError"] < 4 and metric["changedPixelFraction"] < .06, f"independent renderer difference: {name}: {metric}")
        contacts.extend((f"{name} page {i + 1}", page) for i, page in enumerate(poppler))
        report[name] = {"pages": len(reader.pages), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "renderer_differences": metrics}
    encrypted = PdfReader(FIXTURES / "encrypted.pdf")
    require(encrypted.is_encrypted and encrypted.decrypt("wrong-password") == 0, "wrong password was accepted")
    require(encrypted.decrypt("fixture-reader") != 0 and len(encrypted.pages) == 1, "reader password failed")
    settings = encrypted.trailer["/Encrypt"]
    require(encrypted.pages[0].get_contents().get_data() == PdfReader(FIXTURES / "invoice.pdf").pages[0].get_contents().get_data(), "encryption changed visible page commands")
    permissions = int(settings["/P"])
    require(settings["/Length"] == 256 and not permissions & 4 and not permissions & 16, "AES-256/permission settings differ")
    report["encryption"] = {"bits": 256, "wrong_password": "rejected", "print": False, "extract": False}
    encrypted_poppler, encrypted_pdfium = render_encrypted(FIXTURES / "encrypted.pdf", output / "encrypted")
    require(difference(encrypted_poppler, encrypted_pdfium, 1)["meanAbsoluteChannelError"] < 4, "encrypted renderers differ")
    require(difference(encrypted_poppler, output / "invoice/poppler/page-1.png")["meanAbsoluteChannelError"] == 0,
            "encrypted visual output changed")
    contacts.append(("encrypted page 1", encrypted_poppler))
    word = Document(FIXTURES / "shared.docx")
    require(word.tables[0].cell(0, 1).text == "CELL-02 a\u0301 العربية لا 123", "DOCX rich cell text differs")
    with ZipFile(FIXTURES / "shared.docx") as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        require(len(xml.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink")) == 1, "DOCX hyperlink markup missing")
        require(any(r.get("Target") == "https://example.com/docs?a=1&b=2" for r in relationships), "DOCX hyperlink relationship missing")
    report["docx"] = "independently parsed rich cells and hyperlink relationship"
    if args.openssl:
        report["signature"] = verify_signature(FIXTURES / "signed.pdf", args.openssl, output)
        signed_poppler = render_poppler(FIXTURES / "signed.pdf", output / "signed/poppler")
        signed_pdfium = render_pdfium(FIXTURES / "signed.pdf", output / "signed/pdfium")
        require(len(signed_poppler) == len(signed_pdfium) == 1, "signed renderer page counts differ")
        require(difference(signed_poppler[0], signed_pdfium[0], 1)["meanAbsoluteChannelError"] < 4, "signed renderers differ")
        require(difference(signed_poppler[0], output / "invoice/poppler/page-1.png")["meanAbsoluteChannelError"] == 0,
                "signature changed visible content")
        contacts.append(("signed page 1", signed_poppler[0]))
    if args.verapdf_report:
        xml = ET.parse(args.verapdf_report).getroot()
        validation = next(element for element in xml.iter() if element.tag.rsplit("}", 1)[-1] == "validationReport")
        require(validation.get("isCompliant") == "true", "veraPDF rejected archival fixture")
        item = next(element for element in xml.iter() if element.tag.rsplit("}", 1)[-1] == "item")
        require(int(item.get("size")) == (FIXTURES / "archive.pdf").stat().st_size, "veraPDF report describes a different archive size")
        require(args.verapdf_report.stat().st_mtime >= (FIXTURES / "archive.pdf").stat().st_mtime,
                "veraPDF report predates the final archival fixture")
        report["pdfa"] = {**validation.attrib, **list(validation)[0].attrib}
    typography = extracted(ROOT / "examples/build/production/typography.pdf")
    require("OpenType / Unicode" in typography, "invisible text anchor interleaved the typography heading")
    report["reading_order"] = "PDFium extracts the complete typography heading before body text"
    contact_sheet(contacts, output)
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
