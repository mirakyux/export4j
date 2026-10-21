# SPDX-License-Identifier: Apache-2.0
"""Independent Markdown PDF/DOCX output checks and complete two-renderer page verification."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from verify_advanced_pdf import difference, render_pdfium, render_poppler, require
from verify_production import contact_sheet, extracted

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "examples/build/markdown-templates"
OUTPUT = ROOT / "build/verification/markdown-templates"


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = FIXTURES / "summary.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == 1, "Markdown summary should have one page")
    text = extracted(path)
    for value in ("青禾巡检摘要", "Reviewed", "Pending", "Verified", "**数据标记不再解析**。", "已完成现场核查", "无需访问网络"):
        require(value in text, f"Markdown output lost literal content: {value}")
    fonts = {str(font.get_object().get("/BaseFont")) for page in reader.pages
             for font in page["/Resources"]["/Font"].values()}
    require(any("BoldItalic" in f for f in fonts) and any("-Italic" in f for f in fonts)
            and any("-Bold" in f and "BoldItalic" not in f for f in fonts), "real font variants missing")
    images = sum(obj.get_object().get("/Subtype") == "/Image" for page in reader.pages
                 for obj in page["/Resources"].get("/XObject", {}).values())
    require(images == 1 and all(not page.get("/Annots") for page in reader.pages), "image/link policy differs")
    word = Document(FIXTURES / "summary.docx")
    require(word.paragraphs[0].text == "青禾巡检摘要" and word.paragraphs[0].style.name == "Heading 1", "heading semantic lost")
    word_text = "\n".join(p.text for p in word.paragraphs)
    require("**数据标记不再解析**。" in word_text and len(word.inline_shapes) == 1, "DOCX literal binding/image lost")
    styled = {r.text: r.style for p in word.paragraphs for r in p.runs if r.text in {"Reviewed", "Pending", "Verified"}}
    require(styled["Reviewed"].font.bold and styled["Pending"].font.italic
            and styled["Verified"].font.bold and styled["Verified"].font.italic, "DOCX emphasis semantics lost")
    with ZipFile(FIXTURES / "summary.docx") as package:
        require(sum(name.startswith("word/media/") for name in package.namelist()) == 1, "DOCX media inventory differs")
    first = render_poppler(path, OUTPUT / "poppler")
    second = render_pdfium(path, OUTPUT / "pdfium")
    require(len(first) == len(second) == 1, "Markdown renderer page count differs")
    diff = difference(first[0], second[0], 1)
    require(diff["meanAbsoluteChannelError"] < 4 and diff["changedPixelFraction"] < .06, "Markdown renderers differ")
    pointer = ROOT / "build/verification/consumer-build-directory.txt"
    consumer = Path(pointer.read_text(encoding="utf-8").strip()) / "output/consumer-markdown.pdf"
    require(len(PdfReader(consumer).pages) == 1 and "**literal**" in extracted(consumer), "Markdown consumer literal binding failed")
    consumer_first = render_poppler(consumer, OUTPUT / "consumer/poppler")
    consumer_second = render_pdfium(consumer, OUTPUT / "consumer/pdfium")
    require(len(consumer_first) == len(consumer_second) == 1, "Markdown consumer page count differs")
    consumer_diff = difference(consumer_first[0], consumer_second[0], 1)
    require(consumer_diff["meanAbsoluteChannelError"] < 4 and consumer_diff["changedPixelFraction"] < .06, "consumer renderers differ")
    contact_sheet([("Markdown summary", first[0]), ("Markdown consumer", consumer_first[0])], OUTPUT)
    report = {"pages": 1, "consumerPages": 1, "images": images, "fontVariants": sorted(fonts),
              "docxSemantic": "passed", "renderDifference": diff, "consumerRenderDifference": consumer_diff}
    (OUTPUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
