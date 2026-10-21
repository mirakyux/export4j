# SPDX-License-Identifier: Apache-2.0
"""Independently verify XML Each/Rows block cells, final references and native nested DOCX.

Fixed PDF geometry and values are independent of Java manifests. Source/IR manifests only
cross-check provenance and formatting. Render every complete new page twice; never create
or refresh historical visual baselines.
"""
from __future__ import annotations

import argparse
from importlib import metadata
import json
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document as WordDocument
from pypdf import PdfReader
from verify_adr_form import render
from verify_advanced_pdf import difference, require
from verify_block_cells import recover_grid
from verify_inline_images import image_boxes
from verify_multi_floats import pdf_geometry
from verify_production import actual_text, contact_sheet
from verify_semantic_tables import annotations, check_characters, close, page_logical, table_borders

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "examples/build/xml-block-cells"
OUTPUT = ROOT / "build/verification/xml-block-cells"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
OD = "{urn:com.mirakyux.export4j:semantic:1}"
SVG = "{http://www.w3.org/2000/svg}"


def copied(kind, index):
    suffix = f"@{index}" if kind == "each" else f"@row:row-table:{index}"
    return {name: f"{kind}-{name}{suffix}" for name in ("outer", "before", "inner", "after")}


def source_lines(kind, code):
    before = code + " BEFORE INNER=\ufffc EXT=\ufffc"
    after = code + " AFTER P=\ufffc N=\ufffc\nSELF=\ufffc"
    after += " OUT=\ufffc\nXML-LINK \ufffc" if kind == "each" else " ROW-LINK \ufffc"
    return before, after


def final_lines(kind, code, page):
    result = [f"{code} BEFORE INNER={page} EXT=1", f"{code} AFTER P={page} N=5"]
    result += [f"SELF={page} OUT={page}", "XML-LINK \ufffc"] if kind == "each" else [f"SELF={page} ROW-LINK \ufffc"]
    return result


def check_sources():
    nodes = json.loads((FIXTURES / "source.json").read_text(encoding="utf-8"))["nodes"]
    require(len(nodes) == 17 and len({node["id"] for node in nodes}) == 17, "deep source IDs are not unique and complete")
    index = {node["id"]: node for node in nodes}
    expected_order = ["prefix"]
    for kind in ("each", "row"):
        if kind == "row": expected_order.append("row-table")
        for number in range(2):
            ids = copied(kind, number)
            if kind == "each": expected_order.append(ids["outer"])
            expected_order.extend(ids[name] for name in ("before", "inner", "after"))
            parent, row = (ids["outer"], 1) if kind == "each" else ("row-table", number + 1)
            for block, name in enumerate(("before", "inner", "after")):
                require(index[ids[name]]["path"] == [{"tableId": parent, "row": row, "column": 0, "block": block}],
                        "XML copied descendant path confused inner/outer coordinates")
            texts = source_lines(kind, f"{'E' if kind == 'each' else 'R'}{number + 1}")
            for name, text in zip(("before", "after"), texts, strict=True):
                node = index[ids[name]]
                markers = [i for i, char in enumerate(text) if char == "\ufffc"]
                require(node["text"] == text and node["utf16Length"] == len(text.encode("utf-16-le")) // 2,
                        "XML binding flattened or changed original leaf UTF-16 source")
                require(node["fieldIndices"] == (markers if name == "before" else markers[:-1])
                        and node["imageIndices"] == ([] if name == "before" else markers[-1:]), "XML inline source indices differ")
                require((node["size"], node["step"], node["align"])
                        == ((8, 14, "CENTER") if name == "before" else (10, 18, "RIGHT")),
                        "block paragraph did not inherit its cell or preserve explicit override")
            require(index[ids["before"]]["targets"] == [ids["inner"], "prefix"],
                    "lexical clone did not relocate its own inner target or changed external target")
            after_targets = ["", "", ids["before"], ids["outer"]] if kind == "each" else ["", "", ids["before"]]
            require(index[ids["after"]]["targets"] == after_targets,
                    "Each/Rows descendant field target did not follow its own copy")
            inner = index[ids["inner"]]
            require(inner["anchors"] == ([[0, 0, 2, 2], [0, 2, 1, 1], [1, 2, 1, 1], [2, 0, 3, 1]]
                    if kind == "each" else [[0, 0, 1, 2], [0, 1, 1, 1], [1, 1, 1, 1]]),
                    "copied inner grid lost independent horizontal/vertical spans")
            if kind == "each":
                require(index[ids["outer"]]["anchors"] == [[0, 0, 3, 1], [1, 0, 2, 2], [1, 2, 1, 1], [2, 2, 1, 1]],
                        "Each outer source grid changed")
    expected_order.append("end")
    require([node["id"] for node in nodes] == expected_order, "deep bound source order differs")
    require(index["row-table"]["anchors"] == [[0, 0, 2, 1], [1, 0, 2, 1], [2, 0, 2, 1]], "Rows parent grid differs")
    return index


def close_box(actual, expected, message):
    for key, value in zip(("x", "y", "width", "height"), expected, strict=True):
        close(actual[key], value, message + ": " + key)


def check_report():
    path = FIXTURES / "report.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == 5 and not reader.get_fields(), "XML fixed page/form contract differs")
    for page in reader.pages:
        close(page.mediabox.width, 420, "XML page width differs")
        close(page.mediabox.height, 300, "XML page height differs")
    sources = check_sources()
    manifest = json.loads((FIXTURES / "results.json").read_text(encoding="utf-8"))
    require(manifest["pages"] == 5 and len(manifest["tables"]) == 8 and len(manifest["lines"]) == 14,
            "XML IR fragment count differs from the fixed contract")
    expected, wanted_tables = [], []
    for kind in ("each", "row"):
        for number in range(2):
            page = number + (1 if kind == "each" else 3)
            code = f"{'E' if kind == 'each' else 'R'}{number + 1}"
            values = final_lines(kind, code, page + 1)
            expected.extend(values)
            require(page_logical(reader.pages[page]) == values, "final XML fields/source order escaped the atomic copy")
            ids = copied(kind, number)
            wanted_tables.extend([(ids["outer"] if kind == "each" else "row-table", page,
                                   (20, 20, 380, 204 if kind == "each" else 158),
                                   [0, 1, 2] if kind == "each" else [0, number + 1]),
                                  (ids["inner"], page, (26, 70, 380 * 2 / 3 - 12 if kind == "each" else 368,
                                                       84 if kind == "each" else 56),
                                   [0, 1, 2] if kind == "each" else [0, 1])])
    require(actual_text(reader) == expected and not page_logical(reader.pages[0]), "logical XML source was duplicated/reordered")
    require("XML-PREFIX" in (reader.pages[0].extract_text() or "")
            and "XML-END" in (reader.pages[-1].extract_text() or ""), "prefix or following body misplaced")
    for actual, (identity, page, box, rows) in zip(manifest["tables"], wanted_tables, strict=True):
        require(actual["id"] == identity and actual["page"] == page and actual["rows"] == rows,
                "XML connected source group split across pages")
        close_box(actual, box, "fixed inner/outer table geometry differs")

    _, logical, _ = pdf_geometry(reader, 0)
    require(len(logical) == 14 and [line["source"] for line in logical] == expected, "final PDF logical anchor order differs")
    for line, ir in zip(logical, manifest["lines"], strict=True):
        before = " BEFORE " in line["source"]
        kind = "each" if line["page"] in (1, 2) else "row"
        step = 14 if before else 18
        line["y"] -= step * .8
        close(ir["step"], step, "distinct paragraph line heights were lost")
        close(line["y"], (54 if before else (164 if kind == "each" else 136)
                          + (0 if " AFTER " in line["source"] else 18 if "SELF=" in line["source"] else 36)),
              "final leaf paragraph line box y differs")
        require(ir["page"] == line["page"] and ir["logical"] == line["source"], "IR/final logical fields differ")
        for key in ("x", "y", "width"): close(ir[key], line[key], "IR/final PDF line box differs: " + key)
        width = 380 * 2 / 3 - 12 if kind == "each" else 368
        close(line["x"], 26 + (width - line["width"]) / (2 if before else 1), "center/right paragraph alignment differs")
    for identity in {line["id"] for line in manifest["lines"]}:
        values = [line for line in manifest["lines"] if line["id"] == identity]
        original = sources[identity]["text"]
        require(values[0]["start"] == 0 and values[-1]["end"] == len(original)
                and all(a["end"] + 1 == b["start"] and original[a["end"]] == "\n" for a, b in zip(values, values[1:])),
                "XML leaf UTF-16 provenance changed at explicit line boundaries")

    borders = table_borders(reader)
    require(len(borders) == 26 and not [box for box in borders if box["pageIndex"] == 0],
            "atomic group did not skip the short first region or border count differs")
    for page in (1, 2, 3, 4):
        values = [box for box in borders if box["pageIndex"] == page]
        if page < 3:
            width, column = 380 * 2 / 3, (380 * 2 / 3 - 12) / 3
            wanted = [(20, 20, 380, 28), (20, 48, width, 176), (20 + width, 48, 380 / 3, 88),
                      (20 + width, 136, 380 / 3, 88), (26, 70, column * 2, 56),
                      (26 + column * 2, 70, column, 28), (26 + column * 2, 98, column, 28), (26, 126, column * 3, 28)]
        else:
            wanted = [(20, 20, 380, 28), (20, 48, 380, 130), (26, 70, 184, 56),
                      (210, 70, 184, 28), (210, 98, 184, 28)]
        require(len(values) == len(wanted), "copied table border count differs")
        for actual, box in zip(sorted(values, key=lambda b: (b["y"], b["x"])), sorted(wanted, key=lambda b: (b[1], b[0])), strict=True):
            close_box(actual, box, "independently painted merged rectangles differ")
        text = "".join((reader.pages[page].extract_text() or "").split())
        code = f"{'E' if page < 3 else 'R'}{page if page < 3 else page - 2}"
        for suffix in (("-HEADER", "-SIDE-A", "-SIDE-B", "-INNER-V", "-INNER-A", "-INNER-B", "-INNER-END")
                       if page < 3 else ("-INNER-V", "-INNER-A", "-INNER-B")):
            require(text.count(code + suffix) == 1, "source ink lost or duplicated: " + code + suffix)
        if page >= 3: require(text.count("ROWS-HEADER") == 1, "Rows repeated ordinary header missing")

    links, images = annotations(reader), image_boxes(reader)
    require(len(links) == 22 and len(images) == 4, "XML final link/image count differs")
    for page in (1, 2, 3, 4):
        local = [link for link in links if link["pageIndex"] == page]
        internal = [link for link in local if "targetPageIndex" in link]
        wanted_y = [20, 54, 70, 20] if page < 3 else [20, 54, 70]
        require(sorted(link["destinationY"] for link in internal) == sorted(wanted_y), "copied descendant destination y differs")
        require(sum(link["targetPageIndex"] == 0 for link in internal) == 1
                and all(link["targetPageIndex"] == page for link in internal if link["destinationY"] != 20 or link["targetPageIndex"] != 0),
                "copy target escaped to another instance or external target changed")
        label = "each" if page < 3 else "row"
        require({link["uri"] for link in local if "uri" in link}
                == {f"https://example.com/xml-{label}-text", f"https://example.com/xml-{label}-image"}, "XML passive URL changed")
        image = next(value for value in images if value["page"] == page)
        require(image["width"] == image["height"] == 10, "explicit XML image dimensions changed")
        image_link = next(link for link in local if link.get("uri", "").endswith("-image"))
        for key in ("x", "y", "width", "height"): close(image_link[key], image[key], "image hyperlink leaves its painted image")
        require(26 <= image["x"] < image["x"] + 10 <= (26 + 380 * 2 / 3 - 12 if page < 3 else 394)
                and (200 if page < 3 else 154) <= image["y"] < image["y"] + 10 <= (218 if page < 3 else 172),
                "XML image leaves its leaf line box")
    svg = ET.parse(FIXTURES / "report-preview.svg").getroot()
    require([group.get("aria-label") for group in svg.iter(SVG + "g") if group.get("aria-label") in set(expected)] == expected,
            "XML SVG/PDF logical order differs")
    return {"pages": 5, "boundNodes": 17, "copiedParagraphLines": 14, "paintedBorders": len(borders),
            "links": links, "images": images, "characters": check_characters(path, 20)}


def check_package(path, private):
    with ZipFile(path) as package:
        main = ET.fromstring(package.read("word/document.xml"))
        if not private: require(not any(key.startswith(OD) for node in main.iter() for key in node.attrib), "standard input retained private metadata")
        tables = main.findall(f".//{W}tbl")
        require(len(tables) == 2, "XML native nested table count differs")
        require(recover_grid(tables[0]) == ([3800, 3800], [[0, 0, 2, 1], [1, 0, 2, 2]], 1), "XML native outer grid differs")
        require(recover_grid(tables[1]) == ([2453] * 3, [[0, 0, 2, 2], [0, 2, 1, 1], [1, 2, 1, 1]], 1),
                "XML native inner grid/independent spans differ")
        for table, row_index in ((tables[0], 0), (tables[1], 0), (tables[1], 1)):
            row = table.findall(W + "tr")[row_index]
            for run in row.findall(f"{W}tc/{W}p/{W}r"):
                if run.find(W + "t") is not None:
                    require(run.find(f"{W}rPr/{W}color").get(W + "val") == "224488", "native XML cell color was lost")
        cell = tables[0].findall(W + "tr")[1].find(W + "tc")
        blocks = [child for child in cell if child.tag in (W + "p", W + "tbl")]
        require([block.tag for block in blocks] == [W + "p", W + "tbl", W + "p"], "XML native block order was flattened")
        for paragraph, size, line, alignment, before, after in [(blocks[0], 16, 280, "center", 0, 40), (blocks[2], 20, 360, "right", 80, 0)]:
            props = paragraph.find(W + "pPr")
            spacing = props.find(W + "spacing")
            require(props.find(W + "jc").get(W + "val") == alignment
                    and spacing.get(W + "line") == str(line) and spacing.get(W + "before") == str(before)
                    and spacing.get(W + "after") == str(after), "distinct native paragraph formatting changed")
            require(all(node.get(W + "val") == str(size) for node in paragraph.findall(f"{W}r/{W}rPr/{W}sz")), "native run size differs")
        fields = main.findall(f".//{W}fldSimple")
        require([field.get(W + "instr").split()[0] for field in fields] == ["PAGEREF", "PAGE", "NUMPAGES", "PAGEREF"],
                "XML native final field order differs")
        starts = main.findall(f".//{W}bookmarkStart")
        require(len(starts) == len(main.findall(f".//{W}bookmarkEnd")) == 2, "XML nested target bookmark ownership differs")
        names = {node.get(W + "name") for node in starts}
        require(all(field.get(W + "instr").split()[1] in names for field in fields if field.get(W + "instr").startswith("PAGEREF ")),
                "XML native references have dangling bookmarks")
        rels = ET.fromstring(package.read("word/_rels/document.xml.rels"))
        require(sorted(node.get("Target") for node in rels if node.get("TargetMode") == "External")
                == ["https://example.com/xml-word-image", "https://example.com/xml-word-text"], "XML native external relationships differ")
        require(len(main.findall(f".//{W}drawing")) == 1 and sum(name.startswith("word/media/") for name in package.namelist()) == 1,
                "XML native media was lost/duplicated")
    word = WordDocument(path)
    require(len(word.tables) == 1 and len(word.tables[0].cell(1, 0).tables) == 1, "python-docx did not read XML native nested table")
    return {"tables": 2, "fields": 4, "bookmarks": 2, "privateMetadata": private, "orderedParagraphTableParagraph": True}


def check_consumer(directory):
    path = directory / "output/consumer-xml-block-cells.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == 2 and not reader.get_fields(), "isolated XML caller page/form contract differs")
    links, images, borders = annotations(reader), image_boxes(reader), table_borders(reader)
    require(len(links) == 8 and len(images) == 2 and len(borders) == 6, "isolated XML caller content count differs")
    _, logical, _ = pdf_geometry(reader, 18)
    require(len(logical) == 4, "isolated XML final paragraph count differs")
    for page, code in enumerate(("A", "B")):
        close(reader.pages[page].mediabox.width, 320, "caller width differs")
        close(reader.pages[page].mediabox.height, 180, "caller height differs")
        require(page_logical(reader.pages[page]) == [f"CON-XML-BEFORE {code} REF={page + 1}",
                f"CON-XML-AFTER {code} P={page + 1} N=2 R={page + 1}LINK\ufffc"], "caller scoped clone/final fields differ")
        for offset, y in enumerate((56, 110)):
            line = logical[page * 2 + offset]
            close(line["x"], 26, "caller leaf x differs")
            close(line["y"], y, "caller leaf y differs")
        local_borders = sorted((b for b in borders if b["pageIndex"] == page), key=lambda b: b["y"])
        for actual, box in zip(local_borders, [(20, 20, 280, 30), (20, 50, 280, 84), (26, 74, 268, 30)], strict=True):
            close_box(actual, box, "caller independently painted nested geometry differs")
        local = [link for link in links if link["pageIndex"] == page]
        refs = [link for link in local if "targetPageIndex" in link]
        require(len(refs) == 2 and all(link["targetPageIndex"] == page for link in refs)
                and sorted(link["destinationY"] for link in refs) == [56, 110], "caller deep copy references crossed instances")
        require({link["uri"] for link in local if "uri" in link} == {"https://example.com/xml-block", "https://example.com/xml-block-image"},
                "caller passive URLs differ")
        image = next(image for image in images if image["page"] == page)
        require(image["width"] == image["height"] == 8, "caller image dimensions differ")
        image_link = next(link for link in local if link.get("uri", "").endswith("-image"))
        for key in ("x", "y", "width", "height"): close(image_link[key], image[key], "caller image link differs from painted image")
    with ZipFile(directory / "output/consumer-xml-block-cells.docx") as package:
        main = ET.fromstring(package.read("word/document.xml"))
        tables = main.findall(f".//{W}tbl")
        require(len(tables) == 4, "caller native copied nested table count differs")
        for number in range(2):
            outer, inner = tables[number * 2:number * 2 + 2]
            require(outer.get(OD + "id") == f"xml-outer@{number}" and inner.get(OD + "id") == f"xml-inner@{number}",
                    "caller native descendant copy ID differs")
            require(recover_grid(outer) == ([2800, 2800], [[0, 0, 2, 1], [1, 0, 2, 2]], 1)
                    and recover_grid(inner) == ([5360], [[0, 0, 1, 2]], 1), "caller native independent merge grid differs")
        require(len(main.findall(f".//{W}fldSimple")) == 8 and len(main.findall(f".//{W}bookmarkStart")) == 4
                and len(main.findall(f".//{W}drawing")) == 2 and sum(name.startswith("word/media/") for name in package.namelist()) == 1,
                "caller native fields/bookmarks or shared media count differs")
    return {"pages": 2, "links": links, "images": images, "paintedBorders": 6, "nativeNestedDocx": True,
            "characters": check_characters(path, 20)}


def render_checked(path, name, result, contacts):
    first, second = render(path, OUTPUT / name)
    require(len(first) == len(second) == result["pages"], "independent renderer omitted a complete XML page")
    values = [difference(a, b, 2) for a, b in zip(first, second, strict=True)]
    require(all(value["meanAbsoluteChannelError"] < 4 and value["changedPixelFraction"] < .06 for value in values),
            "independent XML PDF renderers disagree")
    result["rendererDifferences"] = values
    contacts.extend((f"{name} page {index + 1}", page) for index, page in enumerate(first))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--consumer-build-directory", type=Path)
    args = parser.parse_args()
    versions = {name: metadata.version(name) for name in ("pypdf", "pypdfium2", "Pillow")}
    require(versions == {"pypdf": "6.10.0", "pypdfium2": "5.13.0", "Pillow": "12.3.0"}, "use the pinned project verification environment")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    report, contacts = {"tools": versions, "report": check_report()}, []
    render_checked(FIXTURES / "report.pdf", "report", report["report"], contacts)
    report["docx"] = check_package(FIXTURES / "blocks.docx", True)
    report["standardDocx"] = check_package(FIXTURES / "blocks-standard.docx", False)
    require((FIXTURES / "docx-read.txt").read_text(encoding="utf-8")
            == "XML_BLOCKS_STRICT_ROUNDTRIP_PASSED\nXML_NATIVE_BLOCKS_STRICT_IMPORT_PASSED\n", "strict XML DOCX read evidence missing")
    directory = args.consumer_build_directory
    if directory is None:
        pointer = ROOT / "build/verification/consumer-build-directory.txt"
        require(pointer.is_file(), "run the isolated caller or pass --consumer-build-directory")
        directory = Path(pointer.read_text(encoding="utf-8").strip())
    if not directory.is_absolute(): directory = ROOT / directory
    report["consumerBuildDirectory"] = str(directory)
    report["consumer"] = check_consumer(directory)
    render_checked(directory / "output/consumer-xml-block-cells.pdf", "consumer", report["consumer"], contacts)
    require(len(contacts) == 7, "complete XML fixture page count differs")
    contact_sheet(contacts, OUTPUT)
    (OUTPUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"pages": len(contacts), "xmlDeepBindings": "passed", "standardNestedDocx": "passed"}, indent=2))


if __name__ == "__main__":
    main()
