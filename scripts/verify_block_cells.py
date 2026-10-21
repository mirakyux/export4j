# SPDX-License-Identifier: Apache-2.0
"""Independently verify ordered atomic block cells, short regions and standard nested DOCX.

Fixed fixture geometry and final field values are asserted from PDF/OOXML. Java IR manifests
only cross-check source provenance. Every complete new PDF page is rendered by Poppler and
the pinned PDFium; no existing baseline is created or updated.
"""
from __future__ import annotations

import argparse
from importlib import metadata
import json
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document as WordDocument
from pypdf import PdfReader
from verify_adr_form import render
from verify_advanced_pdf import difference, require
from verify_inline_images import image_boxes
from verify_multi_floats import pdf_geometry
from verify_production import actual_text, contact_sheet
from verify_semantic_tables import annotations, check_characters, close, page_logical, table_borders

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "examples/build/block-cells"
OUTPUT = ROOT / "build/verification/block-cells"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
OD = "{urn:com.mirakyux.export4j:semantic:1}"
SVG = "{http://www.w3.org/2000/svg}"


def contract(name):
    pages, width = (3, 380) if name == "atomic" else (2, 182)
    prefix = name.upper()
    sources = [prefix + "-P1 REF=\ufffc", prefix + "-P2 P=\ufffc N=\ufffc\nREF=\ufffc ROOT=\ufffc\nLINK \ufffc"]
    final = [prefix + "-P1 REF=2", prefix + f"-P2 P=2 N={pages}", "REF=2 ROOT=2", "LINK \ufffc"]
    return pages, width, sources, final


def check_sources(name, sources):
    data = json.loads((FIXTURES / f"{name}-source.json").read_text(encoding="utf-8"))["paragraphs"]
    require(len(data) == 2, "ordered block paragraphs were flattened, lost or duplicated")
    for index, (actual, text) in enumerate(zip(data, sources, strict=True)):
        require(actual["id"] == f"{name}-p{index + 1}" and actual["text"] == text,
                "leaf paragraph identity or original source differs")
        require(actual["utf16Length"] == len(text.encode("utf-16-le")) // 2, "leaf UTF-16 length differs")
        markers = [i for i, value in enumerate(text) if value == "\ufffc"]
        require(actual["fieldIndices"] == (markers if index == 0 else markers[:-1])
                and actual["imageIndices"] == ([] if index == 0 else markers[-1:]), "leaf field/image source indices differ")
        require(actual["path"] == [{"tableId": f"{name}-outer", "row": 1, "column": 0, "block": index * 2}],
                "block source path confused outer and inner cells")
    return data


def check_fixture(name):
    pages, width, sources, final = contract(name)
    path = FIXTURES / f"{name}.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == pages and not reader.get_fields(), "fixed block-cell page budget/form contract differs")
    for page in reader.pages:
        close(page.mediabox.width, 420, "page width differs")
        close(page.mediabox.height, 300, "page height differs")
    marked = actual_text(reader)
    require(marked == final, "final nested PAGE/PAGES/PAGEREF values or logical source order differ")
    require(page_logical(reader.pages[0]) == [] and page_logical(reader.pages[1]) == final,
            "atomic connected group did not move out of the short remaining region")
    all_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    compact_text = "".join(all_text.split())
    prefix, inner_prefix = name.upper(), name[0].upper() + "I"
    for token in [prefix + "-PREFIX", prefix + "-AFTER", prefix + "-SIDE-A", prefix + "-SIDE-B",
                  *[inner_prefix + suffix for suffix in ("-V", "-A", "-B", "-E")],
                  *[prefix + f"-ROW-{i}" for i in range(1, 6)]]:
        # Outline anchor text can be invisible to ordinary extraction; these tokens are simple ink.
        require(compact_text.count(token) == 1, "real source text was lost/duplicated: " + token)
    require(compact_text.count(prefix + "-HEADER") == 2, "ordinary outer header did not repeat per region")
    require(prefix + "-AFTER" in (reader.pages[-1].extract_text() or ""), "following body paragraph was misplaced")

    provenance = check_sources(name, sources)
    manifest = json.loads((FIXTURES / f"{name}-results.json").read_text(encoding="utf-8"))
    require(manifest["pages"] == pages and len(manifest["tables"]) == 3 and len(manifest["lines"]) == 4,
            "IR source presentations differ from the fixed contract")
    wanted_tables = [(f"{name}-outer", 1, 20, 20, width, 246, list(range(7))),
                     (f"{name}-inner", 1, 23, 61, width * 2 / 3 - 6, 66, [0, 1, 2]),
                     (f"{name}-outer", 2 if name == "atomic" else 1, 20 if name == "atomic" else 218,
                      20, width, 44, [0, 7])]
    for actual, (identity, page, x, y, w, h, rows) in zip(manifest["tables"], wanted_tables, strict=True):
        require(actual["id"] == identity and actual["page"] == page and actual["rows"] == rows,
                "inner/outer grids or atomic source row allocation differ")
        for key, value in {"x": x, "y": y, "width": w, "height": h}.items():
            close(actual[key], value, "fixed table geometry differs: " + key)
    _, anchors, _ = pdf_geometry(reader, 16)
    require(len(anchors) == 4 and [anchor["source"] for anchor in anchors] == final, "final PDF logical anchors differ")
    wanted_y = [45, 127, 143, 159]
    for index, (actual, source_line) in enumerate(zip(anchors, manifest["lines"], strict=True)):
        require(actual["page"] == source_line["page"] == 1 and source_line["logical"] == final[index],
                "final field presentation escaped its atomic cell")
        close(actual["x"], 23, "nested paragraph final PDF x differs")
        # pdf_geometry reports the nonpainting logical anchor's fixed line box, not font ink ascent.
        close(actual["y"], wanted_y[index], "nested paragraph final PDF line box y differs")
        for key in ("x", "y", "width"):
            close(actual[key], source_line[key], "IR/final PDF line geometry differs: " + key)
        require(source_line["id"] == f"{name}-p{1 if index == 0 else 2}", "line provenance lost its paragraph ID")
    require(manifest["lines"][0]["start"] == 0 and manifest["lines"][0]["end"] == len(sources[0]),
            "first leaf UTF-16 range differs")
    p2_lines = manifest["lines"][1:]
    require(p2_lines[0]["start"] == 0 and p2_lines[-1]["end"] == len(sources[1])
            and all(left["end"] + 1 == right["start"] and sources[1][left["end"]] == "\n"
                    for left, right in zip(p2_lines, p2_lines[1:])),
            "last paragraph source ranges were rebased onto another leaf or reordered")

    borders = table_borders(reader)
    require(len(borders) == 14, "inner/outer merge borders were lost or drawn twice")
    large = [box for box in borders if box["pageIndex"] == 1 and abs(box["width"] - width * 2 / 3) < .003
             and abs(box["height"] - 136) < .003]
    require(len(large) == 1, "outer two-row/two-column rectangle is no longer atomic")
    close(large[0]["x"], 20, "outer merged origin x differs")
    close(large[0]["y"], 42, "outer merged origin y differs")
    links, images = annotations(reader), image_boxes(reader)
    require(len(links) == 5 and len(images) == 1 and images[0]["page"] == 1
            and images[0]["width"] == images[0]["height"] == 10, "nested passive links/image count or size differs")
    external = [link for link in links if "uri" in link]
    require({link["uri"] for link in external} == {f"https://example.com/{name}-text", f"https://example.com/{name}-image"},
            "nested external targets changed")
    internal = [link for link in links if "targetPageIndex" in link]
    require(len(internal) == 3 and all(link["pageIndex"] == link["targetPageIndex"] == 1 for link in internal),
            "nested block references did not resolve final physical pages")
    require(sorted(round(link["destinationY"], 3) for link in internal) == [20, 45, 61],
            "table, paragraph and inner-table references lost distinct destinations")
    image_link = next(link for link in external if link["uri"].endswith("-image"))
    for key in ("x", "y", "width", "height"):
        close(image_link[key], images[0][key], "nested image link differs from final painted image")
    require(23 <= images[0]["x"] < images[0]["x"] + 10 <= 20 + width * 2 / 3 - 3
            and 159 <= images[0]["y"] < images[0]["y"] + 10 <= 175, "nested image leaves its outer content box")
    svg = ET.parse(FIXTURES / f"{name}-preview.svg").getroot()
    require(len(svg.findall(f"{SVG}svg")) == pages and len(svg.findall(f".//{SVG}image")) == 1
            and not svg.findall(f".//{SVG}a"), "static SVG page/image contract differs")
    logical = [node.get("aria-label") for node in svg.iter(f"{SVG}g") if node.get("aria-label") is not None]
    require([value for value in logical if value in set(final)] == final, "PDF/SVG nested logical source order differs")
    return {"pages": pages, "paintedBorders": len(borders), "paragraphs": provenance, "links": links,
            "images": images, "characters": check_characters(path, 20), "fixedPdfGeometryAndFinalFields": True}


def recover_grid(table):
    widths = [int(column.get(W + "w")) for column in table.findall(f"{W}tblGrid/{W}gridCol")]
    anchors, active, continuations = [], {}, 0
    rows = table.findall(f"{W}tr")
    for row_index, row in enumerate(rows):
        column, following = 0, {}
        for cell in row.findall(f"{W}tc"):
            span = cell.find(f"{W}tcPr/{W}gridSpan")
            count = 1 if span is None else int(span.get(W + "val"))
            merge = cell.find(f"{W}tcPr/{W}vMerge")
            kind = None if merge is None else merge.get(W + "val", "continue")
            if kind == "continue":
                require(column in active and active[column][2] == count, "native continuation interval changed")
                active[column][3] += 1
                following[column] = active[column]
                require(not cell.findall(f".//{W}t") and not cell.findall(f".//{W}tbl")
                        and not cell.findall(f".//{W}fldSimple") and not cell.findall(f".//{W}drawing"),
                        "native continuation duplicated semantic descendants")
                continuations += 1
            else:
                require(kind in (None, "restart"), "invalid native merge state")
                anchor = [row_index, column, count, 1]
                anchors.append(anchor)
                if kind == "restart": following[column] = anchor
            column += count
        require(column == len(widths), "native row does not cover its own complete grid")
        active = following
    return widths, anchors, continuations


def check_package(path, private):
    with ZipFile(path) as package:
        main = ET.fromstring(package.read("word/document.xml"))
        if not private:
            require(not any(key.startswith(OD) for node in main.iter() for key in node.attrib), "standard fixture retained private metadata")
        tables = main.findall(f".//{W}tbl")
        require(len(tables) == 4, "standard nested tables missing/duplicated")
        outer, inner, tail, tail_inner = tables
        widths, anchors, continuations = recover_grid(outer)
        require(widths == [2533] * 3 and anchors == [[0, 0, 3, 1], [1, 0, 2, 2], [1, 2, 1, 1],
                [2, 2, 1, 1], *[[i, 0, 3, 1] for i in range(3, 8)]] and continuations == 1,
                "standard outer merge grid differs")
        inner_widths, inner_anchors, inner_continuations = recover_grid(inner)
        require(inner_widths == [1649] * 3 and inner_anchors == [[0, 0, 2, 2], [0, 2, 1, 1], [1, 2, 1, 1], [2, 0, 3, 1]]
                and inner_continuations == 1, "standard inner grid did not preserve independent spans/content width")
        cell = outer.findall(f"{W}tr")[1].find(f"{W}tc")
        order = [element.tag for element in cell if element.tag in (W + "p", W + "tbl")]
        require(order == [W + "p", W + "tbl", W + "p"], "standard cell source order was flattened")
        require(len(outer.findall(f"{W}tr/{W}trPr/{W}tblHeader")) == 1
                and not inner.findall(f"{W}tr/{W}trPr/{W}tblHeader"), "native repeated outer header boundary differs")
        tail_cell = tail.find(f"{W}tr/{W}tc")
        tail_blocks = [element for element in tail_cell if element.tag in (W + "p", W + "tbl")]
        require([element.tag for element in tail_blocks] == [W + "p", W + "tbl", W + "p"]
                and not list(tail_blocks[-1]) and not tail_blocks[-1].text, "standard trailing paragraph is missing or carries content")
        if private:
            require(tail_blocks[-1].get(OD + "cellTail") == "true", "own structural tail marker missing")
        require(recover_grid(tail_inner)[1] == [[0, 0, 1, 1]], "tail inner table changed")
        fields = main.findall(f".//{W}fldSimple")
        instructions = [field.get(W + "instr") for field in fields]
        require(len(fields) == 5 and sum(value == "PAGE" for value in instructions) == 1
                and sum(value == "NUMPAGES" for value in instructions) == 1
                and sum(value.startswith("PAGEREF ") for value in instructions) == 3, "nested native fields differ")
        starts, ends = main.findall(f".//{W}bookmarkStart"), main.findall(f".//{W}bookmarkEnd")
        require(len(starts) == len(ends) == 3 and len({node.get(W + "id") for node in starts}) == 3,
                "nested targets/bookmarks lost or duplicated")
        rels = ET.fromstring(package.read("word/_rels/document.xml.rels"))
        require(sorted(node.get("Target") for node in rels if node.get("TargetMode") == "External")
                == ["https://example.com/word-image", "https://example.com/word-text"], "native nested relationships differ")
        require(sum(name.startswith("word/media/") for name in package.namelist()) == 1, "native media resource was duplicated")
    word = WordDocument(path)
    require(len(word.tables) == 2 and len(word.tables[0].cell(1, 0).tables) == 1
            and len(word.tables[0].cell(1, 0).tables[0].rows) == 3, "independent python-docx nested structure read failed")
    return {"tables": 4, "outerAnchors": len(anchors), "innerAnchors": len(inner_anchors),
            "nativeFields": 5, "nativeBookmarks": 3, "vMergeContinuations": 2, "structuralTail": True, "privateMetadata": private}


def check_consumer(directory):
    path = directory / "output/consumer-block-cells.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == 1 and not reader.get_fields(), "isolated caller atomic block-cell page/form contract differs")
    close(reader.pages[0].mediabox.width, 320, "caller page width differs")
    close(reader.pages[0].mediabox.height, 240, "caller page height differs")
    text = reader.pages[0].extract_text() or ""
    for source in ("CON-BLOCK-HEADER", "CON-BLOCK-BEFORE", "CON-INNER-A", "CON-INNER-B", "CON-INNER-C", "CON-INNER-D"):
        require(text.count(source) == 1, "caller real source text lost or duplicated: " + source)
    require(actual_text(reader) == ["CON-BLOCK-AFTER P=1 N=1 R=1LINK\ufffc"], "caller nested final fields or source order differ")
    _, anchors, _ = pdf_geometry(reader, 18)
    require(len(anchors) == 1 and anchors[0]["page"] == 0, "caller final paragraph anchor count differs")
    close(anchors[0]["x"], 22, "caller final paragraph x differs")
    close(anchors[0]["y"], 106, "caller final paragraph line box y differs")
    borders = table_borders(reader)
    require(len(borders) == 5, "caller inner/outer merge border count differs")
    wanted = [(20, 20, 280, 22), (20, 42, 280, 84), (22, 62, 138, 44), (160, 62, 138, 22), (160, 84, 138, 22)]
    recovered = sorted((round(box["x"], 3), round(box["y"], 3), round(box["width"], 3), round(box["height"], 3)) for box in borders)
    require(recovered == sorted(wanted), "caller independently painted nested rectangles differ")
    links, images = annotations(reader), image_boxes(reader)
    require(len(links) == 3 and len(images) == 1 and images[0]["page"] == 0
            and images[0]["width"] == images[0]["height"] == 8, "caller links/image differ")
    require({link.get("uri") for link in links if "uri" in link}
            == {"https://example.com/block-cell", "https://example.com/block-cell-image"}, "caller external targets changed")
    reference = next(link for link in links if "targetPageIndex" in link)
    require(reference["targetPageIndex"] == 0, "caller paragraph reference destination differs")
    close(reference["destinationY"], 44, "caller source paragraph anchor differs")
    image_link = next(link for link in links if link.get("uri", "").endswith("-image"))
    for key in ("x", "y", "width", "height"):
        close(image_link[key], images[0][key], "caller image link differs from actual image")
    require(22 <= images[0]["x"] < images[0]["x"] + 8 <= 298 and 106 <= images[0]["y"] < images[0]["y"] + 8 <= 124,
            "caller image leaves its leaf paragraph")
    with ZipFile(directory / "output/consumer-block-cells.docx") as package:
        main = ET.fromstring(package.read("word/document.xml"))
        tables = main.findall(f".//{W}tbl")
        require(len(tables) == 2, "caller standard nested table count differs")
        require(recover_grid(tables[0]) == ([2800, 2800], [[0, 0, 2, 1], [1, 0, 2, 2]], 1), "caller outer native merge grid differs")
        require(recover_grid(tables[1]) == ([2760, 2760], [[0, 0, 1, 2], [0, 1, 1, 1], [1, 1, 1, 1]], 1),
                "caller inner native merge grid/content width differs")
        cell = tables[0].findall(f"{W}tr")[1].find(f"{W}tc")
        require([element.tag for element in cell if element.tag in (W + "p", W + "tbl")]
                == [W + "p", W + "tbl", W + "p"], "caller standard block order differs")
        fields = main.findall(f".//{W}fldSimple")
        require(len(fields) == 3 and fields[0].get(W + "instr") == "PAGE" and fields[1].get(W + "instr") == "NUMPAGES"
                and fields[2].get(W + "instr").startswith("PAGEREF "), "caller native final fields differ")
        require(len(main.findall(f".//{W}bookmarkStart")) == 1 and len(main.findall(f".//{W}drawing")) == 1
                and sum(name.startswith("word/media/") for name in package.namelist()) == 1, "caller nested bookmark/media ownership differs")
    return {"pages": 1, "paintedBorders": len(borders), "links": links, "images": images,
            "characters": check_characters(path, 20), "nativeNestedDocx": True}


def render_checked(path, name, result, contacts):
    first, second = render(path, OUTPUT / name)
    require(len(first) == len(second) == result["pages"], "independent renderer omitted a complete page")
    metrics = [difference(a, b, 2) for a, b in zip(first, second, strict=True)]
    require(all(value["meanAbsoluteChannelError"] < 4 and value["changedPixelFraction"] < .06 for value in metrics),
            "independent renderers disagree")
    result["rendererDifferences"] = metrics
    contacts.extend((f"{name} page {index + 1}", page) for index, page in enumerate(first))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--consumer-build-directory", type=Path)
    args = parser.parse_args()
    versions = {name: metadata.version(name) for name in ("pypdf", "pypdfium2", "Pillow")}
    require(versions == {"pypdf": "6.10.0", "pypdfium2": "5.13.0", "Pillow": "12.3.0"}, "use the pinned project verification environment")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    report, contacts = {"tools": versions}, []
    for name in ("atomic", "columns"):
        result = check_fixture(name)
        render_checked(FIXTURES / f"{name}.pdf", name, result, contacts)
        report[name] = result
    require((FIXTURES / "docx-read.txt").read_text(encoding="utf-8")
            == "ORDERED_BLOCKS_STRICT_ROUNDTRIP_PASSED\nMETADATA_FREE_NESTING_STRICT_IMPORT_PASSED\n", "strict nested DOCX import evidence missing")
    report["docx"] = check_package(FIXTURES / "blocks.docx", True)
    report["standardDocx"] = check_package(FIXTURES / "blocks-standard.docx", False)
    directory = args.consumer_build_directory
    if directory is None:
        pointer = ROOT / "build/verification/consumer-build-directory.txt"
        require(pointer.is_file(), "run isolated caller or pass --consumer-build-directory")
        directory = Path(pointer.read_text(encoding="utf-8").strip())
    if not directory.is_absolute(): directory = ROOT / directory
    report["consumerBuildDirectory"] = str(directory)
    report["consumer"] = check_consumer(directory)
    render_checked(directory / "output/consumer-block-cells.pdf", "consumer", report["consumer"], contacts)
    require(len(contacts) == 6, "complete-page fixture count differs")
    contact_sheet(contacts, OUTPUT)
    (OUTPUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"pages": len(contacts), "atomicOrderedBlockCells": "passed", "standardNestedDocx": "passed"}, indent=2))


if __name__ == "__main__":
    main()
