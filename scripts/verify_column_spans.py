# SPDX-License-Identifier: Apache-2.0
"""Independently check semantic spans, final PDF geometry, source ranges and native DOCX gridSpan.

Run after :examples:columnSpans and the isolated Maven caller in the pinned project Python.
Fixed source text, widths, page allocation and complete-line ranges are specified here, not
inferred from Java manifests. Both independent renderers process every new complete page.
No baseline or rendering threshold is created or adjusted. --word requires genuine Word
update/save artifacts produced by verify_column_spans_word.ps1 on generated fixtures only.
"""
from __future__ import annotations

import argparse
import hashlib
from importlib import metadata
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document as WordDocument
from pypdf import PdfReader
from verify_adr_form import render
from verify_advanced_pdf import difference, require
from verify_inline_images import image_boxes
from verify_multi_floats import pdf_geometry
from verify_production import actual_text, contact_sheet
from verify_semantic_table_flow import blue_ink
from verify_semantic_tables import annotations, check_characters, close, page_logical, table_borders

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "examples/build/column-spans"
OUTPUT = ROOT / "build/verification/column-spans"
TOLERANCE = .003
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
OD = "{urn:com.mirakyux.export4j:semantic:1}"
SVG = "{http://www.w3.org/2000/svg}"


def labels(prefix, count):
    return [f"{prefix}{number:02}" for number in range(1, count + 1)]


def expected_sources(name):
    if name == "whole-rows":
        result = [(0, 0, 3, "WHOLE P=\ufffc N=\ufffc")]
        for number in range(1, 16):
            span = 2 if number % 2 else 1
            result.append((number, 0, span, f"W{number:02}-A"))
            result.append((number, span, 3 - span, f"W{number:02}-B" +
                           (" P=\ufffc N=\ufffc L \ufffc" if number == 15 else "")))
        return result
    if name == "split-rows":
        return [(0, 0, 3, "SPLIT P=\ufffc N=\ufffc"), (1, 0, 2, "\n".join(labels("S", 22))),
                (1, 2, 1, "\n".join(labels("R", 22)) + " L \ufffc P=\ufffc")]
    return [(0, 0, 3, "COL P=\ufffc N=\ufffc"), (1, 0, 1, "\n".join(labels("C", 16))),
            (1, 1, 2, "\n".join(labels("D", 16)) + " P=\ufffc L \ufffc")]


def check_sources(name):
    values = json.loads((FIXTURES / f"{name}-source.json").read_text(encoding="utf-8"))
    expected = expected_sources(name)
    require(len(values) == len(expected), "source anchors were padded into synthetic cells")
    for value, (row, column, span, source) in zip(values, expected, strict=True):
        require((value["row"], value["column"], value["columnSpan"], value["source"])
                == (row, column, span, source), "source span/text/physical start differs")
        require(value["sourceUtf16Length"] == len(source.encode("utf-16-le")) // 2,
                "bound UTF-16 source length differs")
        objects = [index for index, character in enumerate(source) if character == "\ufffc"]
        image = [] if row == 0 or not objects else objects[-1:]
        if name == "split-rows" and row == 1 and column == 2:
            image = objects[:1]
        # A body cell only has objects when it is the final compact anchor with one image.
        require(value["imageIndices"] == image and value["fieldIndices"] == [index for index in objects if index not in image],
                "binding lost or shifted final-anchor image/field UTF-16 indices")
    return {(row, column): (span, source) for row, column, span, source in expected}


def fixed_contract(name):
    if name == "whole-rows":
        # One header plus eight whole body rows fit into 220 pt; seven remain on page 2.
        return {"pages": 2, "size": (420, 260), "step": 18, "widths": (100, 140, 140),
                "regions": [(0, 20, 20, list(range(1, 9))), (1, 20, 20, list(range(9, 16)))]}
    if name == "split-rows":
        # 24 pt header + 6 pt body padding leaves ten 18 pt complete lines per page.
        return {"pages": 3, "size": (420, 260), "step": 18, "widths": (100, 140, 140),
                "regions": [(0, 20, 20, list(range(1, 11))), (1, 20, 20, list(range(11, 21))),
                            (2, 20, 20, list(range(21, 23)))]}
    # 380 pt content, 16 pt gap -> 182 pt columns. The 4 pt grid padding leaves 174 pt.
    # The 100 pt prefix leaves short 160 pt regions. Seven complete 16 pt body lines fit.
    return {"pages": 2, "size": (420, 300), "step": 16, "widths": (58, 58, 58),
            "regions": [(0, 24, 124, list(range(1, 8))), (0, 222, 124, list(range(8, 15))),
                        (1, 24, 24, list(range(15, 17)))]}


def same_box(left, right):
    return left["pageIndex"] == right["pageIndex"] and all(
        abs(left[key] - right[key]) < TOLERANCE for key in ("x", "y", "width", "height"))


def resolved_line(source, name, row, column, page, pages):
    if row == 0:
        return source.replace("\ufffc", str(page + 1), 1).replace("\ufffc", str(pages), 1)
    if name == "whole-rows" and row == 15 and column == 2:
        return source.replace("\ufffc", str(page + 1), 1).replace("\ufffc", str(pages), 1)
    if name == "split-rows" and column == 2 and source.startswith("R22"):
        return source.rsplit("\ufffc", 1)[0] + str(page + 1)
    if name == "columns" and column == 1 and source.startswith("D16"):
        return source.replace("\ufffc", str(page + 1), 1)
    return source


def check_manifest(name, reader, sources, contract):
    manifest = json.loads((FIXTURES / f"{name}-results.json").read_text(encoding="utf-8"))
    require(manifest["pages"] == contract["pages"] and manifest["bodyPageOffset"] == 0,
            "fixture page count or unexpected contents offset differs")
    placements = manifest["placements"]
    require(len(placements) == len(contract["regions"]), "short region or continuation omitted")
    borders = table_borders(reader)
    _, anchors, _ = pdf_geometry(reader, contract["step"])
    final = {(item["page"], item["source"], round(item["x"], 3), round(item["y"], 3)) for item in anchors}
    collected = {}
    final_anchor = None
    header_text = "WHOLE" if name == "whole-rows" else "SPLIT" if name == "split-rows" else "COL"
    for index, (placement, (page, x, y, numbers)) in enumerate(zip(placements, contract["regions"], strict=True)):
        require(placement["pageIndex"] == page, "table fragment moved to an incorrect physical page")
        for key, expected in (("x", x), ("y", y), ("width", sum(contract["widths"]))):
            close(placement[key], expected, "fixed table fragment geometry differs: " + key)
        if name == "columns":
            require(placement["fragmentIndex"] == index and placement["fragmentCount"] == 3,
                    "nested semantic stream indices/count are incomplete")
        expected_rows = [0, *numbers] if name == "whole-rows" else [0, 1]
        require([row["row"] for row in placement["rows"]] == expected_rows, "original body row order differs")
        expected_height = (len(numbers) + 1) * 24 if name == "whole-rows" else contract["step"] * (len(numbers) + 1) + 12
        close(placement["height"], expected_height, "fixed complete-line table height differs")
        for row_position, row in enumerate(placement["rows"]):
            source_row = row["row"]
            require(row["header"] == (source_row == 0) and row["repeatedHeader"] == (source_row == 0 and index > 0),
                    "full-width repeated header identity differs")
            require(row["continuedRow"] == (source_row == 1 and name != "whole-rows" and index > 0),
                    "complete-row/continued-row flag differs")
            expected_cells = sorted((column, span, source) for (original_row, column), (span, source) in sources.items()
                                    if original_row == source_row)
            require(len(row["cells"]) == len(expected_cells), "covered columns became duplicated content anchors")
            row_y = y + row_position * 24 if name == "whole-rows" else y if source_row == 0 else y + contract["step"] + 6
            row_height = 24 if name == "whole-rows" else contract["step"] + 6 if source_row == 0 else len(numbers) * contract["step"] + 6
            for cell, (column, span, original) in zip(row["cells"], expected_cells, strict=True):
                require((cell["column"], cell["columnSpan"]) == (column, span), "compact anchor resolved using a physical list index")
                expected_box = {"pageIndex": page, "x": x + sum(contract["widths"][:column]), "y": row_y,
                                "width": sum(contract["widths"][column:column + span]), "height": row_height}
                actual_box = {"pageIndex": page, **{key: cell[key] for key in ("x", "y", "width", "height")}}
                require(same_box(actual_box, expected_box), "merged cell geometry ignores covered physical widths")
                require(sum(same_box(value, expected_box) for value in borders) == 1,
                        "final PDF has no unique full-span border at the independent fixed geometry")
                original_lines = original.split("\n")
                line_numbers = [1] if source_row == 0 or name == "whole-rows" else numbers
                require(cell["firstLineIndex"] == line_numbers[0] - 1 and len(cell["lines"]) == len(line_numbers),
                        "complete source line number/count changed during continuation")
                for position, (line, number) in enumerate(zip(cell["lines"], line_numbers, strict=True)):
                    source = original_lines[number - 1]
                    start = sum(len(value) + 1 for value in original_lines[:number - 1])
                    end = start + len(source)
                    require(line["start"] == start and line["end"] in (end, end + int(number < len(original_lines))),
                            "cell-local original UTF-16 range changed")
                    resolved = resolved_line(source, name, source_row, column, page, contract["pages"])
                    require(line["logical"] == resolved, "final page field/image/rich source line differs")
                    close(line["x"], expected_box["x"] + 3, "semantic left padding duplicated by column span")
                    close(line["y"], row_y + 3 + position * contract["step"], "complete baseline spacing differs")
                    require(line["width"] <= expected_box["width"] - 6 + TOLERANCE,
                            "source line overflows its merged inner width")
                    if "\ufffc" in source:
                        require((page, resolved, round(line["x"], 3), round(line["y"], 3)) in final,
                                "final independently painted PDF logical anchor/geometry differs from manifest")
                    collected.setdefault((source_row, column), []).append((number, page, resolved))
                    if source_row > 0 and " L " in source:
                        final_anchor = {"pageIndex": page, **line, "cell": expected_box}
        expected_header = f"{header_text} P={page + 1} N={contract['pages']}"
        require(page_logical(reader.pages[page]).count(expected_header)
                == sum(region[0] == page for region in contract["regions"]),
                "repeated full-span header did not resolve its actual physical page")
    for (row, column), (_, source) in sources.items():
        actual = collected[row, column]
        if row > 0:
            require([number for number, _, _ in actual] == list(range(1, len(source.split("\n")) + 1)),
                    "original complete lines were dropped, repeated or reordered")
    require(final_anchor is not None, "final compact source anchor did not retain linked media")
    return placements, final_anchor, len(borders)


def check_links(reader, name, source):
    prefix = "whole" if name == "whole-rows" else "split" if name == "split-rows" else "columns"
    links = annotations(reader)
    images = image_boxes(reader)
    require(len(links) == 2 and len(images) == 1, "last compact source anchor lost/repeated links or image")
    require({item.get("uri") for item in links} == {f"https://example.com/{prefix}-text", f"https://example.com/{prefix}-image"},
            "last-anchor external targets changed")
    image = images[0]
    require(image["page"] == source["pageIndex"] and image["pixels"] == [8, 8], "linked source image moved or was reencoded")
    expected_size = 10 if name == "columns" else 12
    close(image["width"], expected_size, "inline image width differs")
    close(image["height"], expected_size, "inline image height differs")
    box = source["cell"]
    require(box["x"] + 3 <= image["x"] < image["x"] + image["width"] <= box["x"] + box["width"] - 3 + TOLERANCE,
            "last-source image leaves its actual compact anchor")
    image_link = next(item for item in links if item["uri"].endswith("-image"))
    text_link = next(item for item in links if item["uri"].endswith("-text"))
    for key in ("x", "y", "width", "height"):
        close(image_link[key], image[key], "image link does not match independently painted image: " + key)
    ink = blue_ink(reader)
    require(len(ink) == 1 and ink[0]["pageIndex"] == text_link["pageIndex"] == source["pageIndex"],
            "text link/colored source ink left the final anchor page")
    close(text_link["x"], ink[0]["x"], "text link left differs from actual painted ink")
    close(text_link["x"] + text_link["width"], ink[0]["right"], "text link right differs from actual underline")
    close(text_link["y"], source["y"], "text link line box differs")
    close(text_link["height"], source["height"], "text link complete line height differs")
    return links, images


def check_fixture(name):
    contract = fixed_contract(name)
    path = FIXTURES / f"{name}.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == contract["pages"] and not reader.get_fields(), "fixed PDF page count/static contract differs")
    for page in reader.pages:
        close(page.mediabox.width, contract["size"][0], "PDF physical width differs")
        close(page.mediabox.height, contract["size"][1], "PDF physical height differs")
    sources = check_sources(name)
    placements, last, borders = check_manifest(name, reader, sources, contract)
    content = "\n".join((page.extract_text() or "") + "\n" + "\n".join(page_logical(page)) for page in reader.pages)
    if name == "whole-rows":
        expected = [f"W{number:02}-{side}" for number in range(1, 16) for side in ("A", "B")]
        require(re.findall(r"W\d{2}-[AB]", content) == expected, "whole-row original source order/uniqueness differs")
    else:
        for prefix, count in (("S", 22), ("R", 22)) if name == "split-rows" else (("C", 16), ("D", 16)):
            require(re.findall(rf"{prefix}\d{{2}}", content) == labels(prefix, count),
                    "continued complete source lines lost, repeated or reordered")
    if name == "columns":
        require(content.count("COLUMN-PREFIX") == 1, "short initial region prefix lost or repeated")
    links, images = check_links(reader, name, last)
    svg = ET.parse(FIXTURES / f"{name}-preview.svg").getroot()
    require(len(svg.findall(f"{SVG}svg")) == contract["pages"] and len(svg.findall(f".//{SVG}image")) == 1
            and not svg.findall(f".//{SVG}a"), "static SVG page/image/interaction contract differs")
    logical = [node.get("aria-label") for node in svg.iter(f"{SVG}g") if node.get("aria-label") is not None]
    marked = actual_text(reader)
    require([value for value in logical if value in set(marked)] == marked, "static SVG and final PDF logical order differs")
    return {"pages": contract["pages"], "sourceAnchors": len(sources), "tablePresentations": len(placements),
            "paintedBorders": borders, "lastAnchor": last, "links": links, "images": images,
            "characters": check_characters(path, 20), "fixedGeometryAndOriginalUtf16Checked": True}


def check_package(path, expected_spans, expected_widths, fields, targets):
    with ZipFile(path) as package:
        main = ET.fromstring(package.read("word/document.xml"))
        tables = main.findall(f".//{W}tbl")
        require(len(tables) == 1, "standard DOCX table missing or duplicated")
        table = tables[0]
        widths = [int(value.get(W + "w")) for value in table.findall(f"{W}tblGrid/{W}gridCol")]
        require(widths == expected_widths, "standard DOCX tblGrid physical widths differ")
        rows = table.findall(f"{W}tr")
        spans = [[int(cell.find(f"{W}tcPr/{W}gridSpan").get(W + "val"))
                  if cell.find(f"{W}tcPr/{W}gridSpan") is not None else 1 for cell in row.findall(f"{W}tc")]
                 for row in rows]
        require(spans == expected_spans and all(sum(row) == 3 for row in spans), "standard DOCX gridSpan/compact anchors differ")
        require(not main.findall(f".//{W}vMerge") and len(table.findall(f"{W}tr/{W}trPr/{W}tblHeader")) == 1,
                "DOCX invented vertical spans or lost repeating header")
        native = main.findall(f".//{W}fldSimple")
        require([value.get(W + "instr") for value in native] == fields, "native DOCX field kinds/order differ")
        require(all(float(value.get(OD + "width")) == (18 if len(rows) == 3 else 10) for value in native),
                "explicit native field reservation width differs")
        relationships = ET.fromstring(package.read("word/_rels/document.xml.rels"))
        external = [value.get("Target") for value in relationships if value.get("TargetMode") == "External"]
        require(sorted(external) == sorted(targets), "standard DOCX text/image hyperlink targets differ")
        require(sum(value.startswith("word/media/") for value in package.namelist()) == 1, "standard DOCX inline media missing/duplicated")
        require(not any("vbaProject" in value for value in package.namelist()), "fixture must not contain macros")
    imported = WordDocument(path)
    require(len(imported.tables) == 1 and len(imported.tables[0].rows) == len(expected_spans), "independent python-docx table read failed")
    return {"rows": len(rows), "gridSpan": spans, "gridWidthsTwips": widths, "nativeFields": fields, "externalLinks": external}


def check_consumer(directory):
    path = directory / "output/consumer-column-spans.pdf"
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == 1 and not reader.get_fields(), "isolated caller PDF page/static contract differs")
    close(reader.pages[0].mediabox.width, 320, "caller page width differs")
    close(reader.pages[0].mediabox.height, 220, "caller page height differs")
    expected = [(20, 20, 280, 24), (20, 44, 170, 24), (190, 44, 110, 24),
                (20, 68, 60, 24), (80, 68, 220, 24)]
    borders = table_borders(reader)
    require(len(borders) == len(expected), "caller span borders duplicated/omitted")
    for actual, (x, y, width, height) in zip(borders, expected, strict=True):
        for key, value in (("x", x), ("y", y), ("width", width), ("height", height)):
            close(actual[key], value, "caller independently painted span border differs")
    content = (reader.pages[0].extract_text() or "") + "\n" + "\n".join(page_logical(reader.pages[0]))
    require(all(content.count(value) == 1 for value in ("CON-COLS", "CON-MERGED-A", "CON-B", "CON-MERGED-C"))
            and page_logical(reader.pages[0]) == ["P=1LINK\ufffc"], "caller original/final field source differs")
    links, images = annotations(reader), image_boxes(reader)
    require(len(links) == 2 and len(images) == 1 and images[0]["width"] == images[0]["height"] == 8,
            "caller final compact-anchor media/links differ")
    require({value.get("uri") for value in links} == {"https://example.com/columns", "https://example.com/columns-image"},
            "caller source links changed")
    image_link = next(value for value in links if value["uri"].endswith("-image"))
    for key in ("x", "y", "width", "height"):
        close(image_link[key], images[0][key], "caller image hyperlink differs from independently painted image")
    require(192 <= links[0]["x"] < 298 and 192 <= images[0]["x"] < images[0]["x"] + 8 <= 298,
            "caller final source index was confused with physical column")
    package = check_package(directory / "output/consumer-column-spans.docx", [[3], [2, 1], [1, 2]],
                            [1200, 2200, 2200], ["PAGE"], ["https://example.com/columns", "https://example.com/columns-image"])
    return {"pages": 1, "paintedBorders": len(borders), "links": links, "images": images,
            "characters": check_characters(path, 20), "package": package}


def check_word():
    report = json.loads((FIXTURES / "word-report.json").read_text(encoding="utf-8-sig"))
    require(report["sourceSha256"] == hashlib.sha256((FIXTURES / "spans.docx").read_bytes()).hexdigest(),
            "Word evidence belongs to another generated source")
    require(report["firstUpdateError"] == report["secondUpdateError"] == 0, "Word native field update failed")
    native = [field for field in report["fields"] if field["instruction"].split()[0] != "HYPERLINK"]
    require([field["instruction"] for field in native] == ["PAGE", "NUMPAGES"] and len(native) == 2,
            "Word lost native final-anchor fields")
    pages = report["pages"]
    require(pages >= 2 and native[0]["value"] == str(native[0]["page"]) == str(pages)
            and native[1]["value"] == str(pages), "Word final actual page/count fields differ")
    with ZipFile(FIXTURES / "word-updated.docx") as package:
        main = ET.fromstring(package.read("word/document.xml"))
        table = main.find(f".//{W}tbl")
        require(table is not None, "Word save removed semantic table")
        spans = [[int(cell.find(f"{W}tcPr/{W}gridSpan").get(W + "val"))
                  if cell.find(f"{W}tcPr/{W}gridSpan") is not None else 1 for cell in row.findall(f"{W}tc")]
                 for row in table.findall(f"{W}tr")]
        require(spans == [[3], *[[2, 1] if number % 2 else [1, 2] for number in range(1, 16)]],
                "Word update/save/readback changed compact span geometry")
    independent = WordDocument(FIXTURES / "word-updated.docx")
    require(len(independent.tables) == 1 and len(independent.tables[0].rows) == 16,
            "independent Word-saved table readback differs")
    reader = PdfReader(FIXTURES / "word.pdf", strict=True)
    require(len(reader.pages) == pages, "Word exported PDF page count differs from live reader")
    content = "\n".join(page.extract_text() or "" for page in reader.pages)
    require(re.findall(r"W\d{2}-[AB]", content) == [f"W{number:02}-{side}" for number in range(1, 16) for side in ("A", "B")],
            "Word final PDF original source order/uniqueness differs")
    require(re.sub(r"\s", "", f"W15-B P={pages} N={pages}") in re.sub(r"\s", "", reader.pages[-1].extract_text()),
            "Word final-anchor fields not rendered on their actual page")
    return {"pages": pages, "wordVersion": report["wordVersion"], "nativeFields": native,
            "sourceSha256": report["sourceSha256"], "savedGridSpanAndSourceOrderChecked": True}


def render_checked(path, name, result, contacts):
    first, second = render(path, OUTPUT / name)
    require(len(first) == len(second) == result["pages"], "independent renderer omitted a complete page")
    metrics = [difference(a, b, 2) for a, b in zip(first, second, strict=True)]
    require(all(value["meanAbsoluteChannelError"] < 4 and value["changedPixelFraction"] < .06 for value in metrics),
            "independent renderers disagree")
    result["rendererDifferences"] = metrics
    contacts.extend((f"{name} page {index + 1}", page) for index, page in enumerate(first))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--consumer-build-directory", type=Path)
    parser.add_argument("--word", action="store_true", help="Require genuine generated-fixture Word update/save artifacts")
    args = parser.parse_args()
    versions = {name: metadata.version(name) for name in ("pypdf", "pypdfium2", "Pillow")}
    require(versions == {"pypdf": "6.10.0", "pypdfium2": "5.13.0", "Pillow": "12.3.0"}, "use the pinned project Python")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    report, contacts = {"tools": versions, "renderComparison": {"dpi": 192, "blurPixels": 2, "meanLimit": 4, "changedFractionLimit": .06}}, []
    for name in ("whole-rows", "split-rows", "columns"):
        result = check_fixture(name)
        render_checked(FIXTURES / f"{name}.pdf", name, result, contacts)
        report[name] = result
    require((FIXTURES / "docx-rejected.txt").read_text(encoding="utf-8") == "REPEATED_HEADER_PAGE_REJECTED\n",
            "DOCX repeated-header PAGE rejection missing")
    require(json.loads((FIXTURES / "bindings.json").read_text(encoding="utf-8"))
            == {"javaValueBlock": True, "xmlDynamicRows": True, "compactAnchors": True}, "template binding evidence missing")
    expected_spans = [[3], *[[2, 1] if number % 2 else [1, 2] for number in range(1, 16)]]
    report["package"] = check_package(FIXTURES / "spans.docx", expected_spans, [2000, 2800, 2800],
                                      ["PAGE", "NUMPAGES"], ["https://example.com/whole-text", "https://example.com/whole-image"])
    directory = args.consumer_build_directory
    if directory is None:
        pointer = ROOT / "build/verification/consumer-build-directory.txt"
        require(pointer.is_file(), "run isolated caller or pass --consumer-build-directory")
        directory = Path(pointer.read_text(encoding="utf-8").strip())
    if not directory.is_absolute():
        directory = ROOT / directory
    result = check_consumer(directory)
    render_checked(directory / "output/consumer-column-spans.pdf", "consumer", result, contacts)
    report["consumerBuildDirectory"], report["consumer"] = str(directory), result
    require(len(contacts) == 8, "fixed sample page budget differs")
    if args.word:
        result = check_word()
        render_checked(FIXTURES / "word.pdf", "word", result, contacts)
        report["word"] = result
    contact_sheet(contacts, OUTPUT)
    (OUTPUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"pages": len(contacts), "sourceAndFixedGeometry": "passed", "standardDocxGridSpan": "passed",
                      "word": "passed" if args.word else "not-requested"}, indent=2))


if __name__ == "__main__":
    main()
