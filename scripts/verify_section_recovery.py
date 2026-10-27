# SPDX-License-Identifier: Apache-2.0
"""Independent section, native-band, nested-table, resource and all-page recovery checks."""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import struct
import xml.etree.ElementTree as ET

from docx import Document
from PIL import Image, ImageChops
from pypdf import PdfReader
from verify_adr_form import render
from verify_advanced_pdf import difference, require
from verify_complex_recovery import HtmlText, MARKDOWN, embedded as standard_embedded, unzip
from verify_inline_images import image_boxes
from verify_production import contact_sheet
from verify_semantic_tables import page_logical

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "examples/build/section-recovery"
OUTPUT = ROOT / "build/verification/section-recovery"
FORMATS = {"PDF": "pdf", "DOCX": "docx", "HTML": "html", "MARKDOWN": "md", "SVG": "svg", "PNG_ZIP": "png.zip"}
EXACT = {"PDF", "SVG", "PNG_ZIP"}
WEB = {"HTML", "MARKDOWN"}
BODIES = ["BODY-A1 ${literal}", "BODY-A2", "BODY-A3", "BODY-B1"]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
S = "{http://www.w3.org/2000/svg}"
BLUE, PURPLE = (34, 85, 170), (136, 34, 68)


def text(node):
    return "".join(node.itertext())


def embedded(data, key):
    if key != "PDF": return standard_embedded(data, key)
    reader = PdfReader(io.BytesIO(data), strict=True)
    require(len(reader.pages) == 4 and set(reader.attachments) == {"export4j-recovery.zip"}, "PDF attachment/pages differ")
    specification = reader.trailer["/Root"]["/Names"]["/EmbeddedFiles"]["/Names"][1].get_object()
    match = re.fullmatch(r"Export4j recovery;source-bytes=(\d+)", specification["/Desc"])
    require(match is not None, "PDF source prefix marker missing")
    length = int(match[1])
    require(8 <= length < len(data) and data[:length].startswith(b"%PDF-"), "PDF original prefix differs")
    values = reader.attachments["export4j-recovery.zip"]
    require(len(values) == 1, "PDF attachment is ambiguous")
    return data[:length], values[0]


def resources(entries):
    for name, source in (("fonts/Regular.ttf", "LiberationSans-Regular.ttf"),
                         ("fonts/Bold.ttf", "LiberationSans-Bold.ttf"),
                         ("licenses/Font.txt", "LiberationSans-OFL.txt")):
        require(entries[name] == (ROOT / "assets/fonts" / source).read_bytes(), "explicit font/license changed")
    require(all(name in {"template.xml", "fonts/Regular.ttf", "fonts/Bold.ttf", "licenses/Font.txt"}
                or name.startswith("images/") for name in entries), "unexpected retained source")


def recovery(encoded, standard, key, expected):
    entries = unzip(encoded)
    manifest = ET.fromstring(entries.pop("recovery.xml"))
    require(manifest.get("version") == "1" and manifest.get("format") == key, "manifest format/version differs")
    require(int(manifest.get("bytes")) == len(standard) and manifest.get("sha256") == hashlib.sha256(standard).hexdigest(), "standard binding differs")
    digest = hashlib.sha256()
    for name, data in sorted(entries.items()):
        name = name.encode("utf-8")
        digest.update(struct.pack(">II", len(name), len(data))); digest.update(name); digest.update(data)
    require(digest.hexdigest() == manifest.get("model-sha256"), "model/resource binding differs")
    resources(entries)
    model = ET.fromstring(entries["template.xml"])
    contents = text(model)
    for token in BODIES + ["BEFORE-INNER", "INNER-HEAD-A", "INNER-HEAD-B", "INNER-" + expected,
                           "INNER-RIGHT", "AFTER-INNER", "OUTER-HEAD-A", "OUTER-HEAD-B", "TALL-CELL", "OUTER-TAIL"]:
        require(contents.count(token) == 1, "model missing/duplicated content: " + key + ": " + token)
    require("UNUSED-REGION-SOURCE" not in contents and (expected == "ORIGINAL" or "ORIGINAL" not in contents), "old source returned")
    require(model.get("literal-text") == "true", "literal text contract lost")
    images = [name for name in entries if name.startswith("images/")]
    require(len(images) == (2 if key in EXACT else 0), "removed/current image inventory differs")
    require(len(model.findall(".//table")) == (1 if key == "MARKDOWN" else 2), "nested table preservation differs")
    if key not in WEB:
        sections = model.findall("section")
        require([s.get("first-page-number") for s in sections] == ["2", "8"], "section numbering lost")
        require([s.get("id") for s in sections] == ["a", "b"], "section IDs changed")
        require(model.find(".//cell[@row-span='2']") is not None, "rowspan lost")
        ids = [n.get("id") for n in model.iter() if n.get("id")]
        require(len(ids) == len(set(ids)), "semantic IDs duplicated")
        definitions = {kind: {} for kind in ("header", "footer", "background", "foreground")}
        for section in sections:
            for kind, known in definitions.items():
                nodes = section.findall(kind)
                for node in nodes:
                    for source in (list(node) if kind in {"background", "foreground"} else [node]):
                        if source.get("source-id"): known[source.get("source-id")] = source
                        else: require(source.get("source-ref") in known, "shared source reference missing")
        for kind in ("header", "footer"):
            nodes = sections[1].findall(kind + "[@variant='first']")
            require(len(nodes) == 1, "explicit empty FIRST pair lost")
            node = nodes[0]
            resolved = definitions[kind].get(node.get("source-ref"), node)
            require(not list(resolved) and resolved.get("height") == "0.0", "empty FIRST became inherited")
        if key in EXACT:
            boxes = sections[0].findall("background/box")
            require([(b.get("x"), b.get("y"), b.get("width")) for b in boxes] == [("24.0", "340.0", "80.0"), ("114.0", "340.0", "80.0")], "background geometry changed")
            require(boxes[1].get("source-ref") == boxes[0].get("source-id"), "overlay shared source copied")
            require(len(sections[0].findall("foreground/box")) == 1, "foreground missing")
        else:
            require(not model.findall(".//box") and not model.findall(".//image"), "DOCX restored discarded layers")
    else:
        require(len(model.findall("section")) == 1 and not model.findall(".//image") and not model.findall(".//page-number"), "Web restored richer source")
    issues = manifest.findall("issue")
    if key in EXACT:
        require(not issues, "lossless target reported losses")
    else:
        require(any(i.get("feature") == "MODEL_RECOVERY" and i.find("effect").get("value") == "CURRENT_FILE_REIMPORT" for i in issues), "actual reimport missing")
        require(any(i.get("stage") == "EXPORT" for i in issues), "export history lost")
        if key in WEB: require(any(i.get("stage") == "IMPORT" for i in issues), "import history lost")
        if key == "HTML": require(any(i.get("feature") == "TABLE_CELL_CONTAINER" for i in issues), "nested container loss missing")
    return entries, manifest


def docx(data, expected):
    entries = unzip(data)
    root, settings = ET.fromstring(entries["word/document.xml"]), ET.fromstring(entries["word/settings.xml"])
    require(settings.find(W + "evenAndOddHeaders") is not None, "DOCX even selector lost")
    doc = Document(io.BytesIO(data))
    require(len(doc.sections) == 2 and len(doc.tables) == 1 and len(doc.tables[0].cell(1, 0).tables) == 1, "native nested structure differs")
    require(len(root.findall(".//" + W + "tbl")) == 2, "native table count differs")
    contents = "".join(n.text or "" for n in root.iter(W + "t"))
    for token in BODIES + ["BEFORE-INNER", "INNER-" + expected, "AFTER-INNER", "TALL-CELL", "OUTER-TAIL"]:
        require(contents.count(token) == 1, "native missing/duplicated content: " + token)
    positions = [contents.index(t) for t in ("BEFORE-INNER", "INNER-" + expected, "AFTER-INNER", "OUTER-TAIL")]
    require(positions == sorted(positions), "native nested-cell order differs")
    require([n.get(W + "val", "continue") for n in root.iter(W + "vMerge")] == ["restart", "continue"], "native rowspan differs")
    require(all(len(grid.findall(W + "gridCol")) == 2 for grid in root.iter(W + "tblGrid")), "native grid differs")
    rels = {n.get("Id"): "word/" + n.get("Target") for n in ET.fromstring(entries["word/_rels/document.xml.rels"])}
    sections = list(root.iter(W + "sectPr"))
    require(len(sections) == 2, "native section count differs")
    for index, section in enumerate(sections):
        require(section.find(W + "pgNumType").get(W + "start") == str([2, 8][index]), "native restart differs")
        require(section.find(W + "titlePg") is not None, "native explicit FIRST lost")
        size = section.find(W + "pgSz")
        require((size.get(W + "w"), size.get(W + "h")) == ("8400", "8400"), "native paper geometry differs")
        for kind in ("header", "footer"):
            bands = {n.get(W + "type"): ET.fromstring(entries[rels[n.get(R + "id")]]) for n in section.findall(W + kind + "Reference")}
            require(set(bands) == {"default", "first", "even"}, "native variant references differ")
            for variant, band in bands.items():
                label = {"default": "SHARED-" + expected if kind == "header" else "FOOT-BAND",
                         "first": "FIRST-BAND" if kind == "header" and index == 0 else "",
                         "even": "EVEN-BAND" if kind == "header" else "EVEN-FOOT"}[variant]
                require("".join(n.text or "" for n in band.iter(W + "t")) == (label + " PAGE=/" if label else ""), "native band text differs")
                require([n.get(W + "instr").strip() for n in band.iter(W + "fldSimple")] == (["PAGE", "NUMPAGES"] if label else []), "native band fields differ")
    require(not any(n.startswith("word/media/") for n in entries), "DOCX hidden overlay resource returned")


def pixels(image, scale=1):
    for x, expected in ((30, BLUE), (50, PURPLE), (90, BLUE), (120, BLUE)):
        actual = image.convert("RGB").getpixel((round(x * scale), round(348 * scale)))
        require(max(abs(a - b) for a, b in zip(actual, expected, strict=True)) <= 1, "overlay location/color/order differs")


def standard(data, key, expected):
    if key == "DOCX": docx(data, expected)
    elif key == "PNG_ZIP":
        pages = unzip(data)
        require(set(pages) == {f"page-{i:05d}.png" for i in range(1, 5)}, "PNG page inventory differs")
        for value in pages.values():
            with Image.open(io.BytesIO(value)) as image:
                image.load(); require(image.size == (420, 420) and image.format == "PNG", "PNG geometry differs"); pixels(image)
    elif key == "SVG":
        root = ET.fromstring(data)
        pages = root.findall(S + "svg")
        require(len(pages) == 4 and root.get("viewBox") == "0 0 420.0 1740.0", "stacked SVG pages differ")
        for index, page in enumerate(pages):
            require(page.get("data-page") == str(index + 1) and float(page.get("y")) == index * 440
                    and page.get("viewBox") == "0 0 420.0 420.0", "SVG per-page geometry differs")
            require(len(page.findall(".//" + S + "image")) == 3, "SVG layer count differs")
    elif key in WEB:
        html = data.decode("utf-8") if key == "HTML" else MARKDOWN.render(data.decode("utf-8"))
        contents = "".join(HtmlText(html).text)
        for token in BODIES + ["BEFORE-INNER", "INNER-" + expected, "AFTER-INNER", "TALL-CELL", "OUTER-TAIL"]:
            require(contents.count(token) == 1, "standard Web content missing/duplicated: " + token)
        require(html.count("<table") == (2 if key == "HTML" else 1), "standard Web table structure differs")


def pdf(path, label, key, expected, contacts):
    reader = PdfReader(path, strict=True)
    pages = 1 if key in WEB else 4
    require(len(reader.pages) == pages, "PDF page count differs: " + label)
    contents = "\n".join(page.extract_text() or "" for page in reader.pages)
    for token in BODIES + ["BEFORE-INNER", "INNER-" + expected, "AFTER-INNER", "TALL-CELL", "OUTER-TAIL"]:
        require(contents.count(token) == 1, "PDF missing/duplicated content: " + label + ": " + token)
    if expected != "ORIGINAL": require("ORIGINAL" not in contents and not any("ORIGINAL" in t for p in reader.pages for t in page_logical(p)), "old text returned")
    positions = [contents.index(t) for t in ["BEFORE-INNER", "INNER-" + expected, "AFTER-INNER", "OUTER-TAIL"]]
    require(positions == sorted(positions), "nested content order differs")
    boxes = image_boxes(reader)
    for index, page in enumerate(reader.pages):
        if key not in WEB:
            require((float(page.mediabox.width), float(page.mediabox.height)) == (420, 420), "PDF geometry differs")
            own = page.extract_text() or ""
            require([body for body in BODIES if body in own] == [BODIES[index]], "PDF body page differs")
            labels = [t for t in page_logical(page) if "PAGE=" in t]
            header = ["FIRST-BAND", "SHARED-" + expected, "EVEN-BAND", ""][index]
            footer = ["", "FOOT-BAND", "EVEN-FOOT", ""][index]
            require(labels == [s + f" PAGE={[2, 3, 4, 8][index]}/4" for s in (header, footer) if s], "final field/variant differs: " + label)
        actual = [b for b in boxes if b["page"] == index]
        require(len(actual) == (3 if key in EXACT else 0), "PDF image count differs")
        if actual:
            require([(b["x"], b["width"]) for b in actual] == [(24, 80), (114, 80), (44, 40)], "PDF layer paint order/width differs")
    left, right = render(path, OUTPUT / label)
    metrics = []
    for index, (a, b) in enumerate(zip(left, right, strict=True)):
        delta = difference(a, b, blur_radius=2)
        require(delta["meanAbsoluteChannelError"] < 4 and delta["changedPixelFraction"] < .06, "independent render mismatch: " + label)
        if key in EXACT:
            for path in (a, b):
                with Image.open(path) as image: pixels(image, 192 / 72)
        metrics.append(delta); contacts.append((f"{label}-{index + 1}", a))
    return left, metrics


def exact_pixels(a, b):
    with Image.open(a) as x, Image.open(b) as y:
        return ImageChops.difference(x.convert("RGB"), y.convert("RGB")).getbbox() is None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--consumer-build-directory", type=Path, required=True)
    args = parser.parse_args()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    contacts, reports, controls, edited = [], {}, {}, {}
    original_xml = None
    for key, ext in FORMATS.items():
        directory, models = SOURCE / key.lower(), {}
        for name in ("original", "cycle-3", "edited") + (() if key in EXACT else ("ignored",)):
            expected = "EDITED" if name == "edited" else "ORIGINAL"
            data = (directory / f"{name}.{ext}").read_bytes()
            package = (directory / f"{name}.{ext}.recovery.zip").read_bytes()
            entries, manifest = recovery(package, data, key, expected)
            models[name] = entries["template.xml"]
            plain, payload = embedded((directory / f"embedded-{name}.{ext}").read_bytes(), key)
            require(plain == data and payload == package, "embedding changed accurate standard/model bytes")
            standard(data, key, expected)
            if key in EXACT:
                if original_xml is None: original_xml = entries["template.xml"]
                if name != "edited": require(entries["template.xml"] == original_xml, "exact model/cycle differs")
        control = unzip((directory / "cycle-3.xml.zip").read_bytes())
        require(control["template.xml"] == models["original"], "XML control changed model")
        if key not in EXACT: require(models["original"] == models["ignored"], "WARN/IGNORE model differs")
        controls[key], reports[key.lower() + "-control"] = pdf(directory / "model-control.pdf", key.lower() + "-control", key, "ORIGINAL", contacts)
        edited[key], reports[key.lower() + "-edited"] = pdf(directory / "model-edited.pdf", key.lower() + "-edited", key, "EDITED", contacts)
    base, reports["source"] = pdf(SOURCE / "source.pdf", "source", "PDF", "ORIGINAL", contacts)
    for key in EXACT:
        require(all(exact_pixels(a, b) for a, b in zip(base, controls[key], strict=True)), "unchanged exact target changed pixels")
        require([exact_pixels(a, b) for a, b in zip(base, edited[key], strict=True)] == [False, False, True, True], "edit escaped nested/shared selected pages")
    native = SOURCE / "docx/native.docx"
    docx(native.read_bytes(), "ORIGINAL")
    _, reports["native-docx"] = pdf(SOURCE / "native-docx.pdf", "native-docx", "DOCX", "ORIGINAL", contacts)
    matrix = list(csv.DictReader((SOURCE / "matrix.csv").open(encoding="utf-8")))
    require(len(matrix) == 36 and {(r["a"], r["b"]) for r in matrix} == {(a, b) for a in FORMATS for b in FORMATS}, "matrix incomplete")
    for row in matrix:
        require(int(row["first_losses"]) <= int(row["b_losses"]) <= int(row["final_losses"]), "matrix dropped actual history")
        require((row["exact_model"] == "true") == (row["a"] in EXACT and row["b"] in EXACT), "matrix boundary differs")
    policies = list(csv.DictReader((SOURCE / "policy.csv").open(encoding="utf-8")))
    for key in set(FORMATS) - EXACT:
        rows = {r["action"]: r for r in policies if r["format"] == key}
        require(set(rows) == {"WARN", "IGNORE", "ERROR"}, "policy action coverage missing")
        require(rows["WARN"]["total"] == rows["IGNORE"]["total"] and int(rows["WARN"]["warnings"]) > 0 and rows["IGNORE"]["warnings"] == "0", "WARN/IGNORE history differs")
    caller = args.consumer_build_directory / "output"
    for key, ext in (("PNG_ZIP", "png.zip"), ("DOCX", "docx")):
        plain, payload = embedded((caller / ("consumer-section." + ext)).read_bytes(), key)
        recovery(payload, plain, key, "CALLER"); standard(plain, key, "CALLER")
    _, reports["consumer"] = pdf(caller / "consumer-section.pdf", "consumer", "PDF", "CALLER", contacts)
    contact_sheet(contacts, OUTPUT)
    (OUTPUT / "report.json").write_text(json.dumps({"matrix": 36, "exact_pairs": 9, "policies": policies,
        "pdf_pages": len(contacts), "renders": reports}, indent=2), encoding="utf-8")
    print(f"Verified six-format section/variant/layer/nested-table recovery, 36 A-B-A cases and {len(contacts)} PDF pages with two renderers")


if __name__ == "__main__":
    main()
