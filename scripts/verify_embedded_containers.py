# SPDX-License-Identifier: Apache-2.0
"""Inspect optional PDF, DOCX and PNG page ZIP recovery with independent readers."""
import argparse
import base64
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import struct
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from PIL import Image
from docx import Document as WordDocument
from pypdf import PdfReader, PdfWriter
from verify_advanced_pdf import difference, render_pdfium, render_poppler, require
from verify_production import check_bundle

ROOT = Path(__file__).resolve().parents[1]
SOURCE_ZIP = ROOT / "examples/build/embedded-zip-recovery"
SOURCE_PDF = ROOT / "examples/build/embedded-pdf-recovery"
OUTPUT = ROOT / "build/verification/embedded-containers"
PARTS = {"DOCX": ("docx", "docx", "customXml/export4j-recovery.xml"),
         "PNG_ZIP": ("png_zip", "png.zip", "export4j/recovery.zip")}


def bundle(encoded, standard, format_key, expected, binding=True):
    with ZipFile(io.BytesIO(encoded)) as archive:
        require(archive.testzip() is None and len(archive.namelist()) == len(set(archive.namelist())), "recovery ZIP CRC/duplicate failure")
        entries = {name: archive.read(name) for name in archive.namelist()}
    manifest = ET.fromstring(entries.pop("recovery.xml"))
    require(manifest.tag == "recovery" and manifest.get("version") == "1" and manifest.get("format") == format_key, "recovery version/format differs")
    matches = int(manifest.get("bytes")) == len(standard) and manifest.get("sha256") == hashlib.sha256(standard).hexdigest()
    require(matches == binding, "standard-file binding differs")
    digest = hashlib.sha256()
    for name, content in sorted(entries.items()):
        key = name.encode("utf-8")
        digest.update(struct.pack(">II", len(key), len(content)))
        digest.update(key)
        digest.update(content)
    require(digest.hexdigest() == manifest.get("model-sha256"), "model/resource digest differs")
    require(set(entries) == {"template.xml", "fonts/Regular.ttf", "licenses/Font.txt"}, "unexpected hidden source/resource")
    require(entries["fonts/Regular.ttf"] == (ROOT / "assets/fonts/LiberationSans-Regular.ttf").read_bytes(), "embedded font differs")
    require(entries["licenses/Font.txt"] == (ROOT / "assets/fonts/LiberationSans-OFL.txt").read_bytes(), "embedded license differs")
    xml = entries["template.xml"].decode("utf-8")
    words = [word for paragraph in ET.fromstring(xml).findall(".//p") for word in "".join(paragraph.itertext()).split()]
    require(words == [expected, "EMBEDDED-PDF" if format_key == "PDF" else "EMBEDDED-ZIP", "UNCHANGED-CONTROL"], "saved model differs")
    require("UNUSED-SOURCE" not in xml and (expected == "ORIGINAL" or "ORIGINAL" not in xml), "old source/model survived")
    return entries["template.xml"]


def zip_projection(data, part):
    with ZipFile(io.BytesIO(data)) as archive:
        require(archive.testzip() is None and len(archive.namelist()) == len(set(archive.namelist())), "container ZIP CRC/duplicate failure")
        matches = [info for info in archive.infolist() if info.filename == part]
        require(len(matches) == 1 and matches[0].compress_type == 0, "recovery ZIP member missing/duplicated/compressed")
        info, payload = matches[0], archive.read(part)
    footer = data.rfind(b"PK\x05\x06")
    require(footer >= 0 and footer + 22 + struct.unpack_from("<H", data, footer + 20)[0] == len(data), "ZIP end record differs")
    count, size, central = struct.unpack_from("<HII", data, footer + 10)
    require(central + size == footer, "ZIP central directory bounds differ")
    offset, found, record_size = central, None, None
    while offset < footer:
        require(data[offset:offset + 4] == b"PK\x01\x02", "invalid central member")
        name_size, extra_size, comment_size = struct.unpack_from("<HHH", data, offset + 28)
        end = offset + 46 + name_size + extra_size + comment_size
        require(end <= footer, "central member length differs")
        if data[offset + 46:offset + 46 + name_size].decode("utf-8") == part:
            require(found is None, "duplicate recovery central member")
            found, record_size = offset, end - offset
        offset = end
    require(offset == footer and found is not None and found + record_size == footer, "recovery central record is not last")
    start = info.header_offset
    name_size, extra_size = struct.unpack_from("<HH", data, start + 26)
    require(start + 30 + name_size + extra_size + info.compress_size == central, "recovery local record is not last")
    end = bytearray(data[footer:]); struct.pack_into("<HHII", end, 8, count - 1, count - 1, size - record_size, start)
    return data[:start] + data[central:found] + end, payload


def pdf_attachment(data):
    reader = PdfReader(io.BytesIO(data))
    require(len(reader.pages) == 1 and set(reader.attachments) == {"export4j-recovery.zip"}, "PDF attachment/page differs")
    specification = reader.trailer["/Root"]["/Names"]["/EmbeddedFiles"]["/Names"][1].get_object()
    match = re.fullmatch(r"Export4j recovery;source-bytes=(\d+)", specification["/Desc"])
    require(match is not None, "PDF source length marker missing")
    length = int(match.group(1))
    require(8 <= length < len(data) and data[:length].startswith(b"%PDF-"), "PDF original prefix differs")
    values = reader.attachments["export4j-recovery.zip"]
    require(len(values) == 1, "PDF attachment is ambiguous")
    return data[:length], values[0]


def inspect_zip(path, format_key, expected):
    folder, extension, part = PARTS[format_key]
    data = path.read_bytes()
    standard, payload = zip_projection(data, part)
    if format_key == "DOCX":
        xml = ET.fromstring(payload)
        require(xml.tag == "export4j-recovery" and xml.get("version") == "1", "DOCX custom XML differs")
        payload = base64.b64decode(xml.text, validate=True)
        require([p.text for p in WordDocument(io.BytesIO(data)).paragraphs] == [expected + " EMBEDDED-ZIP", "UNCHANGED-CONTROL"], "Word text/order differs")
        with ZipFile(io.BytesIO(data)) as archive:
            types = ET.fromstring(archive.read("[Content_Types].xml"))
            require(any(item.get("Extension") == "xml" and item.get("ContentType") == "application/xml" for item in types), "DOCX recovery part content type missing")
    else:
        with ZipFile(io.BytesIO(data)) as archive:
            require([name for name in archive.namelist() if name != part] == ["page-00001.png"], "PNG page ZIP member differs")
            with Image.open(io.BytesIO(archive.read("page-00001.png"))) as image:
                require(image.format == "PNG" and image.size == (794, 1123), "PNG page dimensions differ")
    model = bundle(payload, standard, format_key, expected)
    return standard, model


def inspect_pdf(path, expected):
    data = path.read_bytes()
    standard, encoded = pdf_attachment(data)
    reader = PdfReader(io.BytesIO(data))
    require(reader.pages[0].extract_text().split() == [expected, "EMBEDDED-PDF", "UNCHANGED-CONTROL"], "PDF visible text/order differs")
    model = bundle(encoded, standard, "PDF", expected)
    return standard, model


def rewrite_zip(data, replacements, omit=()):
    result = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as original, ZipFile(result, "w") as changed:
        for info in original.infolist():
            if info.filename not in omit:
                changed.writestr(info, replacements.get(info.filename, original.read(info.filename)))
    return result.getvalue()


def prepare_external():
    for format_key, (folder, extension, part) in PARTS.items():
        directory = SOURCE_ZIP / folder
        original = (directory / f"original.{extension}").read_bytes()
        if format_key == "DOCX":
            with ZipFile(io.BytesIO(original)) as archive:
                document = archive.read("word/document.xml")
            require(document.count(b"ORIGINAL") == 1, "DOCX external edit target differs")
            modified = rewrite_zip(original, {"word/document.xml": document.replace(b"ORIGINAL", b"EXTERNAL")})
        else:
            with ZipFile(io.BytesIO(original)) as archive:
                image = archive.read("page-00001.png")
            with Image.open(io.BytesIO(image)) as page:
                page = page.convert("RGB"); page.paste((0, 80, 255), (10, 10, 30, 30))
                stream = io.BytesIO(); page.save(stream, "PNG")
            modified = rewrite_zip(original, {"page-00001.png": stream.getvalue()})
        (directory / f"external.{extension}").write_bytes(modified)
        (directory / f"stripped.{extension}").write_bytes(rewrite_zip(original, {}, omit=(part,)))
    reader = PdfReader(SOURCE_PDF / "original.pdf")
    writer = PdfWriter(); writer.clone_document_from_reader(reader)
    writer.add_metadata({"/Title": "EXTERNAL PDF EDIT"})
    writer.pages[0].cropbox.upper_right = (float(reader.pages[0].mediabox.width) - 20, float(reader.pages[0].mediabox.height))
    writer.write(SOURCE_PDF / "external.pdf")
    writer = PdfWriter()
    for page in reader.pages: writer.add_page(page)
    writer.write(SOURCE_PDF / "stripped.pdf")
    print("External ZIP/PDF rewrites prepared with current document/page changes and removed recovery")


def pdf_render(path, label, expected):
    reader = PdfReader(path)
    require(len(reader.pages) == 1 and reader.pages[0].extract_text().split() == [expected, "EMBEDDED-PDF" if label.startswith("pdf-") else "EMBEDDED-ZIP", "UNCHANGED-CONTROL"], "PDF export text differs")
    poppler = render_poppler(path, OUTPUT / label / "poppler")
    pdfium = render_pdfium(path, OUTPUT / label / "pdfium")
    require(len(poppler) == len(pdfium) == 1, "PDF renderer page count differs")
    comparison = difference(poppler[0], pdfium[0], blur_radius=1)
    require(comparison["meanAbsoluteChannelError"] < 4 and comparison["changedPixelFraction"] < .06, "PDF renderer output differs")
    return comparison


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prepare-external", action="store_true")
    parser.add_argument("--consumer-build-directory", type=Path)
    args = parser.parse_args()
    if args.prepare_external:
        prepare_external(); return
    require(args.consumer_build_directory is not None, "provide the current isolated consumer directory")
    consumer = args.consumer_build_directory.resolve()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    sizes, renders = {}, {}
    for format_key, (folder, extension, part) in PARTS.items():
        directory = SOURCE_ZIP / folder
        plain = (directory / f"plain.{extension}").read_bytes()
        original, model = inspect_zip(directory / f"original.{extension}", format_key, "ORIGINAL")
        require(original == plain, "embedding changed original ZIP bytes")
        cycled, cycle_model = inspect_zip(directory / f"cycle-3.{extension}", format_key, "ORIGINAL")
        require(cycled == original and cycle_model == model and (directory / f"cycle-3.{extension}").read_bytes() == (directory / f"original.{extension}").read_bytes(), "ZIP three-round exact control differs")
        inspect_zip(directory / f"edited.{extension}", format_key, "EDITED")
        inspect_zip(consumer / f"output/consumer-embedded-{format_key.lower()}.{extension}", format_key, "CALLER")
        with ZipFile(directory / f"external.{extension}") as archive:
            require(part in archive.namelist(), "external ZIP did not retain stale recovery")
            if format_key == "DOCX":
                require("EXTERNAL" in WordDocument(io.BytesIO((directory / f"external.{extension}").read_bytes())).paragraphs[0].text, "external Word content did not change")
            else:
                with Image.open(io.BytesIO(archive.read("page-00001.png"))) as image:
                    require(image.convert("RGB").getpixel((10, 10)) == (0, 80, 255), "external PNG pixel did not change")
        with ZipFile(directory / f"stripped.{extension}") as archive:
            require(part not in archive.namelist(), "stripped ZIP retained recovery")
        sizes[format_key] = {"plain": len(plain), "embedded": (directory / f"original.{extension}").stat().st_size}
        for name in ("control", "edited"):
            renders[f"{folder}-{name}"] = pdf_render(directory / f"{name}.pdf", f"{folder}-{name}", "ORIGINAL" if name == "control" else "EDITED")
    original_pdf, model = inspect_pdf(SOURCE_PDF / "original.pdf", "ORIGINAL")
    require(original_pdf == (SOURCE_PDF / "plain.pdf").read_bytes(), "PDF embedded prefix differs from plain PDF")
    inspect_pdf(SOURCE_PDF / "cycle-3.pdf", "ORIGINAL")
    inspect_pdf(SOURCE_PDF / "edited.pdf", "EDITED")
    inspect_pdf(consumer / "output/consumer-embedded-pdf.pdf", "CALLER")
    require(set(PdfReader(SOURCE_PDF / "external.pdf").attachments) == {"export4j-recovery.zip"}, "external PDF dropped recovery")
    require("export4j-recovery.zip" not in PdfReader(SOURCE_PDF / "stripped.pdf").attachments, "stripped PDF retained recovery")
    sizes["PDF"] = {"plain": len(original_pdf), "embedded": (SOURCE_PDF / "original.pdf").stat().st_size}
    for name in ("original", "cycle-3", "edited"):
        renders[f"pdf-{name}"] = pdf_render(SOURCE_PDF / f"{name}.pdf", f"pdf-{name}", "EDITED" if name == "edited" else "ORIGINAL")
    renders["pdf-caller"] = pdf_render(consumer / "output/consumer-embedded-pdf.pdf", "pdf-caller", "CALLER")
    with (SOURCE_ZIP / "external.csv").open(encoding="utf-8", newline="") as source: zip_rows = list(csv.DictReader(source))
    with (SOURCE_PDF / "external.csv").open(encoding="utf-8", newline="") as source: pdf_rows = list(csv.DictReader(source))
    require(len(zip_rows) == 12 and len(pdf_rows) == 6 and all(row["recovered"] == "false" for row in zip_rows + pdf_rows), "external action matrix incomplete")
    for format_key, (plain, embedded) in {
        "HTML": (ROOT / "examples/build/embedded-html-recovery/plain.html", ROOT / "examples/build/embedded-html-recovery/original.html"),
        "MARKDOWN": (ROOT / "examples/build/embedded-comment-recovery/md/plain.md", ROOT / "examples/build/embedded-comment-recovery/md/original.md"),
        "SVG": (ROOT / "examples/build/embedded-comment-recovery/svg/plain.svg", ROOT / "examples/build/embedded-comment-recovery/svg/original.svg"),
        "PNG": (ROOT / "examples/build/embedded-png-recovery/plain.png", ROOT / "examples/build/embedded-png-recovery/original.png"),
    }.items():
        sizes[format_key] = {"plain": plain.stat().st_size, "embedded": embedded.stat().st_size}
    for size in sizes.values():
        size["added"] = size["embedded"] - size["plain"]
        size["ratio"] = round(size["embedded"] / size["plain"], 2)
        require(size["added"] > 0, "optional recovery did not increase file size")
    require((SOURCE_ZIP / "docx/word-report-current.json").exists(), "run the separate Microsoft Word open check")
    word_report = json.loads((SOURCE_ZIP / "docx/word-report-current.json").read_text(encoding="utf-8-sig"))
    require(word_report["sourceSha256"] == hashlib.sha256((SOURCE_ZIP / "docx/original.docx").read_bytes()).hexdigest(), "Word open does not match this run's DOCX")
    report = {"formats": list(sizes), "sizes": sizes, "externalActions": len(zip_rows) + len(pdf_rows), "pdfPages": len(renders),
              "renders": renders, "wordOpen": "passed; open/save round trip not verified", "localBundle": check_bundle(consumer)}
    (OUTPUT / "report.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report))


if __name__ == "__main__":
    main()
