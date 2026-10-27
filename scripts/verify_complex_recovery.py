# SPDX-License-Identifier: Apache-2.0
"""Independent complex structures, actual-loss reconciliation, embedded bytes and PDF renders."""
import argparse
import base64
import csv
import hashlib
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re
import struct
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document
from markdown_it import MarkdownIt
from PIL import Image, ImageChops
from pypdf import PdfReader
from verify_adr_form import render
from verify_advanced_pdf import difference, require
from verify_embedded_containers import pdf_attachment, zip_projection
from verify_embedded_png_recovery import chunks
from verify_production import contact_sheet

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "examples/build/complex-recovery"
OUTPUT = ROOT / "build/verification/complex-recovery"
FORMATS = {"PDF": "pdf", "DOCX": "docx", "HTML": "html", "MARKDOWN": "md", "SVG": "svg", "PNG": "png", "PNG_ZIP": "png.zip"}
WEB = {"HTML", "MARKDOWN"}
MARKDOWN = MarkdownIt("commonmark").enable("table")


class HtmlText(HTMLParser):
    def __init__(self, source):
        super().__init__()
        self.text = []
        self.feed(source)

    def handle_data(self, data):
        self.text.append(data)


def unzip(data):
    with ZipFile(io.BytesIO(data)) as archive:
        require(archive.testzip() is None and len(archive.namelist()) == len(set(archive.namelist())), "invalid ZIP CRC/inventory")
        return {name: archive.read(name) for name in archive.namelist()}


def recovery(encoded, standard, key, expected):
    entries = unzip(encoded)
    manifest = ET.fromstring(entries.pop("recovery.xml"))
    require(manifest.get("version") == "1" and manifest.get("format") == key, "manifest format/version differs")
    require(int(manifest.get("bytes")) == len(standard) and manifest.get("sha256") == hashlib.sha256(standard).hexdigest(), "standard binding differs")
    digest = hashlib.sha256()
    for name, data in sorted(entries.items()):
        name = name.encode("utf-8")
        digest.update(struct.pack(">II", len(name), len(data))); digest.update(name); digest.update(data)
    require(digest.hexdigest() == manifest.get("model-sha256"), "model/resource binding differs")
    require(entries["fonts/Regular.ttf"] == (ROOT / "assets/fonts/LiberationSans-Regular.ttf").read_bytes(), "regular font changed")
    require(entries["fonts/Bold.ttf"] == (ROOT / "assets/fonts/LiberationSans-Bold.ttf").read_bytes(), "bold font changed")
    require(entries["licenses/Font.txt"] == (ROOT / "assets/fonts/LiberationSans-OFL.txt").read_bytes(), "license changed")
    require(all(name in {"template.xml", "fonts/Regular.ttf", "fonts/Bold.ttf", "licenses/Font.txt"} or name.startswith("images/") for name in entries), "unexpected retained source")
    model = ET.fromstring(entries["template.xml"])
    text = "".join(model.itertext())
    for token in ("COMPLEX-" + expected, "CELL-" + expected, "BODY-CONTROL ${literal}", "HEADER-A", "HEADER-B", "MERGED-CELL", "CELL-UNCHANGED", "BULLET-CONTROL", "NUMBER-CONTROL", "REFERENCE-CONTROL="):
        require(text.count(token) == 1, "complex model missing/duplicated token: " + token)
    require("UNUSED-OLD-SOURCE" not in text and (expected == "ORIGINAL" or "ORIGINAL" not in text), "removed text returned")
    require(len(model.findall(".//table")) == 1 and len(model.findall(".//heading")) == 1 and len(model.findall(".//bullet")) == 1 and len(model.findall(".//numbered")) == 1, "complex model structures differ")
    return entries, manifest, model


def embedded(data, key):
    if key == "PDF":
        return pdf_attachment(data)
    if key in {"DOCX", "PNG_ZIP"}:
        plain, payload = zip_projection(data, "customXml/export4j-recovery.xml" if key == "DOCX" else "export4j/recovery.zip")
        if key == "DOCX":
            xml = ET.fromstring(payload)
            require(xml.tag == "export4j-recovery" and xml.get("version") == "1", "DOCX envelope differs")
            payload = base64.b64decode(xml.text, validate=True)
        return plain, payload
    if key == "PNG":
        parts = chunks(data)
        payloads = [value for name, value, _ in parts if name == b"exJM"]
        require(len(payloads) == 1, "PNG recovery chunk differs")
        return b"\x89PNG\r\n\x1a\n" + b"".join(raw for name, _, raw in parts if name != b"exJM"), payloads[0]
    match = re.fullmatch(rb"(.*)\n<!--export4j-model-recovery\n([A-Za-z0-9+/]*={0,2})\n-->\n", data, re.DOTALL)
    require(match is not None, "comment envelope missing")
    return match[1], base64.b64decode(match[2], validate=True)


def png(data):
    with Image.open(io.BytesIO(data)) as image:
        image.load()
        require(image.format == "PNG" and image.size == (596, 842), "PNG dimensions differ")
        require(image.convert("RGB").getextrema()[0][0] < 64, "PNG page is blank")


def structure(data, key, expected, picture):
    if key == "DOCX":
        doc = Document(io.BytesIO(data))
        require(len(doc.tables) == 1 and len(doc.inline_shapes) == int(picture), "DOCX table/image differs")
        require("COMPLEX-" + expected in "\n".join(p.text for p in doc.paragraphs), "DOCX heading missing")
        require("CELL-" + expected in " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells), "DOCX cell missing")
    elif key == "SVG":
        root = ET.fromstring(data)
        require(root.find(".//{http://www.w3.org/2000/svg}path") is not None, "SVG outlines missing")
        require(len(root.findall(".//{http://www.w3.org/2000/svg}image")) == int(picture), "SVG image differs")
    elif key == "PNG":
        png(data)
    elif key == "PNG_ZIP":
        pages = unzip(data); require(set(pages) == {"page-00001.png"}, "PNG page ZIP inventory differs")
        png(pages["page-00001.png"])
    elif key in WEB:
        html = data.decode("utf-8") if key == "HTML" else MARKDOWN.render(data.decode("utf-8"))
        require(all(tag in html for tag in ("<h1", "<ul", "<ol", "<table")), "Web complex structures missing")
        require(("<img" in html) == picture, "Web image projection differs")
        text = "".join(HtmlText(html).text)
        require("CELL-" + expected in text and "MERGED-CELL" in text, "Web cell text missing")


def pdf(path, label, expected):
    reader = PdfReader(path, strict=True)
    require(len(reader.pages) == 1, "complex PDF should fit one page")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for token in ("COMPLEX-" + expected, "CELL-" + expected, "CELL-UNCHANGED", "MERGED-CELL", "BULLET-CONTROL", "NUMBER-CONTROL"):
        require(text.count(token) == 1, "PDF missing/duplicated content: " + label + ": " + token)
    require(any(re.fullmatch(r"1\.\s*NUMBER-CONTROL", line.strip()) for line in text.splitlines()), "number marker duplicated into body: " + label)
    require(any(re.fullmatch(r"[-•]\s*BULLET-CONTROL", line.strip()) for line in text.splitlines()), "bullet marker duplicated into body: " + label)
    poppler, pdfium = render(path, OUTPUT / label)
    result = difference(poppler[0], pdfium[0], blur_radius=2)
    # Use the established complex XML/table tolerances (verify_xml_semantic_saving), rather than
    # the simple-paragraph preset: Poppler/PDFium antialias thin shared table edges differently.
    require(result["meanAbsoluteChannelError"] < 4 and result["changedPixelFraction"] < .06, "PDF renderers diverge: " + label)
    return poppler[0], result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--consumer-build-directory", type=Path, required=True)
    args = parser.parse_args()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    renders, results = [], {}
    exact_model = None
    exact_renders = {}
    for key, extension in FORMATS.items():
        directory = SOURCE / key.lower()
        models = {}
        for name in ("original", "cycle-3", "edited") + (("ignored",) if key in WEB else ()):
            expected = "EDITED" if name == "edited" else "ORIGINAL"
            data = (directory / f"{name}.{extension}").read_bytes()
            package = (directory / f"{name}.{extension}.recovery.zip").read_bytes()
            entries, manifest, model = recovery(package, data, key, expected)
            models[name] = entries["template.xml"]
            plain, payload = embedded((directory / f"embedded-{name}.{extension}").read_bytes(), key)
            require(plain == data and payload == package, "embedding changed standard bytes/model: " + key)
            structure(data, key, expected, key not in WEB or name in {"original", "ignored"})
            images = [name for name in entries if name.startswith("images/")]
            require(len(images) == (0 if key in WEB else 1), "discarded/current image inventory differs")
            if key not in WEB:
                require(not manifest.findall("issue"), "lossless complex control reported loss")
                require(model.find(".//cell[@column-span='2']") is not None, "merged cell lost")
                if name == "original":
                    if exact_model is None: exact_model = entries["template.xml"]
                    require(entries["template.xml"] == exact_model, "lossless format changed model")
            else:
                issues = manifest.findall("issue")
                require(any(issue.get("feature") == "MODEL_RECOVERY" and issue.find("effect").get("value") == "CURRENT_FILE_REIMPORT" for issue in issues), "reimport loss missing")
                require(any(issue.get("stage") == "IMPORT" and issue.get("feature") == "IMAGE_RESOURCE" for issue in issues), "actual image import loss missing")
                require(not images and model.find(".//image") is None and model.find(".//number-reference") is None, "richer source content returned")
        if key not in WEB:
            require(models["original"] == models["cycle-3"], "lossless target cycles changed model")
        else:
            require(models["original"] == models["ignored"], "WARN and IGNORE projected models differ")
        control = (directory / "cycle-3.xml.zip").read_bytes()
        entries, manifest, _ = recovery((directory / "cycle-3.xml.zip.recovery.zip").read_bytes(), control, "XML_TEMPLATE", "ORIGINAL")
        require(entries["template.xml"] == models["original"], "three XML cycles changed model")
        for name, expected in (("model-control", "ORIGINAL"), ("model-edited", "EDITED")):
            label = key.lower() + "-" + name
            page, result = pdf(directory / (name + ".pdf"), label, expected)
            if key not in WEB:
                if name not in exact_renders: exact_renders[name] = page
                with Image.open(exact_renders[name]) as baseline, Image.open(page) as current:
                    require(ImageChops.difference(baseline.convert("RGB"), current.convert("RGB")).getbbox() is None, "lossless format changed control/edit appearance")
            renders.append((label, page)); results[label] = result
        if key in {"PDF", "DOCX", "HTML", "MARKDOWN"}:
            data = (directory / f"reference.{extension}").read_bytes()
            entries, _, model = recovery((directory / f"reference.{extension}.recovery.zip").read_bytes(), data, key, "ORIGINAL")
            require((model.find(".//number-reference") is not None) == (key not in WEB), "native reference preservation/loss differs")
            plain, payload = embedded((directory / f"embedded-reference.{extension}").read_bytes(), key)
            require(plain == data and payload == (directory / f"reference.{extension}.recovery.zip").read_bytes(), "reference embedding changed bytes")
            if key == "DOCX":
                word = ET.fromstring(unzip(data)["word/document.xml"])
                require(any(node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr", "").startswith("REF ")
                            for node in word.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldSimple")), "DOCX native reference missing")
            path = directory / "reference-model.pdf"
            if key not in WEB:
                require(bool(PdfReader(path).pages[0].get("/Annots")), "native reference destination missing")
            page, result = pdf(path, key.lower() + "-reference", "ORIGINAL")
            renders.append((key.lower() + "-reference", page)); results[key.lower() + "-reference"] = result
    matrix = list(csv.DictReader((SOURCE / "matrix.csv").open(encoding="utf-8")))
    require(len(matrix) == 49 and {(row["a"], row["b"]) for row in matrix} == {(a, b) for a in FORMATS for b in FORMATS}, "complex matrix incomplete")
    for row in matrix:
        require(int(row["first_losses"]) <= int(row["b_losses"]) <= int(row["final_losses"]), "matrix dropped history")
        require((row["exact_model"] == "true") == (row["a"] not in WEB and row["b"] not in WEB), "matrix exact-model boundary differs")
    policies = list(csv.DictReader((SOURCE / "policy.csv").open(encoding="utf-8")))
    for key in WEB:
        rows = {row["action"]: row for row in policies if row["format"] == key}
        require(set(rows) == {"ERROR", "WARN", "IGNORE"}, "action coverage incomplete")
        require(rows["WARN"]["total"] == rows["IGNORE"]["total"] and int(rows["WARN"]["import_losses"]) > 0, "actual losses differ by action")
        require(int(rows["WARN"]["warnings"]) > 0 and rows["IGNORE"]["warnings"] == "0", "warning presentation differs")
    caller = args.consumer_build_directory / "output"
    for action in ("WARN", "IGNORE"):
        data = (caller / f"consumer-lossy-{action}.md").read_bytes()
        standard, package = embedded(data, "MARKDOWN")
        recovery(package, standard, "MARKDOWN", "ORIGINAL")
        standard, package = embedded((caller / f"consumer-lossy-{action}.docx").read_bytes(), "DOCX")
        entries, _, _ = recovery(package, standard, "DOCX", "CALLER")
        require(not any(name.startswith("images/") for name in entries), "caller restored removed images")
        page, result = pdf(caller / f"consumer-lossy-{action}.pdf", "consumer-" + action.lower(), "CALLER")
        renders.append(("consumer-" + action.lower(), page)); results["consumer-" + action.lower()] = result
    contact_sheet(renders, OUTPUT)
    (OUTPUT / "report.json").write_text(json.dumps({"matrix": 49, "exact_model_pairs": 25, "renders": results}, indent=2), encoding="utf-8")
    print("Verified complex seven-format structures, 49 A-B-A cases, native-reference boundaries, actual lossy reimport, published caller and 20 PDF pages with two renderers")


if __name__ == "__main__":
    main()
